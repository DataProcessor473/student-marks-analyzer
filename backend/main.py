from fastapi import FastAPI, HTTPException, Query, Depends, UploadFile, File, Request, BackgroundTasks, WebSocket, WebSocketDisconnect
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, FileResponse, StreamingResponse
from fastapi.security import OAuth2PasswordBearer
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, validator
from typing import List, Optional, Dict, Any
from jose import JWTError, jwt
from datetime import datetime, timedelta
from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.util import get_remote_address
from slowapi.errors import RateLimitExceeded
from apscheduler.schedulers.background import BackgroundScheduler
from apscheduler.triggers.interval import IntervalTrigger
import bcrypt
import numpy as np
import sqlite3
import os
import io
import re
import random
import string
import smtplib
import csv as csv_module
import secrets
import json
import pyotp
import qrcode
import shutil
import zipfile
import asyncio
from pathlib import Path
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from contextlib import contextmanager
from dotenv import load_dotenv

load_dotenv()

try:
    from pdf_generator import generate_report_card
except ImportError:
    generate_report_card = None
    print("WARNING: pdf_generator.py not found - classic PDF disabled")

try:
    from pdf_templates import generate_modern_report, generate_minimal_report
except ImportError:
    generate_modern_report = None
    generate_minimal_report = None
    print("WARNING: pdf_templates.py not found - modern PDF disabled")

try:
    from models_phase3 import (
        StudentProfileUpdate, ExamCreate, TimetableCreate,
        AssignmentCreate, SubmissionUpdate,
        FeeStructureCreate, FeePaymentCreate,
        SavedFilterCreate, ScheduledReportCreate,
        WhatsAppOTPRequest, BackupRestoreRequest
    )
except ImportError:
    print("WARNING: models_phase3.py not found")


# ============================================================
# CONFIG
# ============================================================
SECRET_KEY = os.getenv("SECRET_KEY", secrets.token_urlsafe(48))
REFRESH_SECRET_KEY = os.getenv("REFRESH_SECRET_KEY", secrets.token_urlsafe(48))
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = int(os.getenv("ACCESS_TOKEN_EXPIRE_MINUTES", 30))
REFRESH_TOKEN_EXPIRE_DAYS = 7
MAX_FAILED_ATTEMPTS = 5
LOCKOUT_DURATION_MINUTES = 15

OTP_EXPIRE_MINUTES = int(os.getenv("OTP_EXPIRE_MINUTES", 10))
OTP_LENGTH = int(os.getenv("OTP_LENGTH", 6))

SMTP_HOST = os.getenv("SMTP_HOST", "smtp.gmail.com")
SMTP_PORT = int(os.getenv("SMTP_PORT", 587))
SMTP_USER = os.getenv("SMTP_USER", "").strip()
SMTP_PASS = os.getenv("SMTP_PASS", "").strip()
SMTP_FROM = os.getenv("SMTP_FROM", SMTP_USER).strip()

DEV_MODE_SMS = os.getenv("DEV_MODE_SMS", "true").lower() == "true"
FAST2SMS_API_KEY = os.getenv("FAST2SMS_API_KEY", "").strip()

TOTP_ISSUER = "StudentMarksAnalyzer"

# Phase 3
UPLOAD_DIR = os.getenv("UPLOAD_DIR", "uploads")
MAX_UPLOAD_SIZE_MB = int(os.getenv("MAX_UPLOAD_SIZE_MB", 5))
BACKUP_DIR = os.getenv("BACKUP_DIR", "backups")
AUTO_BACKUP_HOURS = int(os.getenv("AUTO_BACKUP_HOURS", 24))
WEBSOCKET_ENABLED = os.getenv("WEBSOCKET_ENABLED", "true").lower() == "true"

TWILIO_ACCOUNT_SID = os.getenv("TWILIO_ACCOUNT_SID", "")
TWILIO_AUTH_TOKEN = os.getenv("TWILIO_AUTH_TOKEN", "")
TWILIO_WHATSAPP_FROM = os.getenv("TWILIO_WHATSAPP_FROM", "whatsapp:+14155238886")

REPORTS_DIR = os.path.join(os.path.dirname(__file__), "reports")
os.makedirs(REPORTS_DIR, exist_ok=True)
os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(BACKUP_DIR, exist_ok=True)
os.makedirs(os.path.join(UPLOAD_DIR, "photos"), exist_ok=True)
os.makedirs(os.path.join(BACKUP_DIR, "auto"), exist_ok=True)

try:
    from twilio.rest import Client as TwilioClient
    TWILIO_AVAILABLE = True
except ImportError:
    TWILIO_AVAILABLE = False
    print("WARNING: twilio not installed - WhatsApp disabled")


# ============================================================
# APP
# ============================================================
app = FastAPI(
    title="Student Marks Analyzer API",
    description="Complete API with all features + Phase 3",
    version="11.0.0",
)

limiter = Limiter(key_func=get_remote_address)
app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

oauth2_scheme = OAuth2PasswordBearer(tokenUrl="/auth/login", auto_error=False)


# ============================================================
# VALIDATION
# ============================================================
def validate_password_strength(password: str) -> tuple:
    if len(password) < 8:
        return False, "Password must be at least 8 characters"
    if not re.search(r"[A-Z]", password):
        return False, "Must contain an uppercase letter"
    if not re.search(r"[a-z]", password):
        return False, "Must contain a lowercase letter"
    if not re.search(r"\d", password):
        return False, "Must contain a digit"
    if not re.search(r"[!@#$%^&*(),.?\":{}|<>_\-+=\[\]\\\/;'`~]", password):
        return False, "Must contain a special character"
    return True, ""


def validate_subject_name(name: str) -> tuple:
    if not name or not name.strip():
        return False, "Subject name cannot be empty"
    name = name.strip()
    if len(name) < 2:
        return False, "Subject name too short"
    if len(name) > 50:
        return False, "Subject name too long"
    if not name[0].isalpha():
        return False, "Subject name must start with a letter"
    if not re.search(r"[aeiouAEIOU]", name):
        return False, f"'{name}' doesn't look like a valid subject"
    if not re.match(r"^[A-Za-z][A-Za-z0-9\s\-\.\(\)&,']*$", name):
        return False, "Invalid characters in subject name"
    letters = re.sub(r"[^A-Za-z]", "", name).lower()
    if letters and len(set(letters)) <= 2:
        return False, "Subject name has no variety"
    return True, ""


def validate_email(email: str) -> bool:
    return bool(re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email))


def validate_phone(phone: str) -> tuple:
    if not phone:
        return False, "Phone required"
    cleaned = re.sub(r"[\s\-\(\)]", "", phone)
    if not re.match(r"^\+?[0-9]{10,15}$", cleaned):
        return False, "Invalid phone number"
    if not cleaned.startswith("+"):
        cleaned = "+91" + cleaned.lstrip("0") if len(cleaned) == 10 else "+" + cleaned
    return True, cleaned


def verify_password(plain: str, hashed: str) -> bool:
    try:
        pwd_bytes = plain.encode("utf-8")[:72]
        hash_bytes = hashed.encode("utf-8") if isinstance(hashed, str) else hashed
        return bcrypt.checkpw(pwd_bytes, hash_bytes)
    except Exception:
        return False


def get_password_hash(password: str) -> str:
    pwd_bytes = password.encode("utf-8")[:72]
    return bcrypt.hashpw(pwd_bytes, bcrypt.gensalt(rounds=12)).decode("utf-8")


# ============================================================
# TOKENS
# ============================================================
def create_access_token(data: dict) -> str:
    d = data.copy()
    d["type"] = "access"
    d["exp"] = datetime.utcnow() + timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
    return jwt.encode(d, SECRET_KEY, algorithm=ALGORITHM)


def create_refresh_token(data: dict) -> str:
    d = data.copy()
    d["type"] = "refresh"
    d["exp"] = datetime.utcnow() + timedelta(days=REFRESH_TOKEN_EXPIRE_DAYS)
    return jwt.encode(d, REFRESH_SECRET_KEY, algorithm=ALGORITHM)


def decode_token(token: str, refresh: bool = False) -> Optional[dict]:
    try:
        key = REFRESH_SECRET_KEY if refresh else SECRET_KEY
        return jwt.decode(token, key, algorithms=[ALGORITHM])
    except JWTError:
        return None


# ============================================================
# EMAIL & OTP
# ============================================================
def generate_otp() -> str:
    return "".join(random.choices(string.digits, k=OTP_LENGTH))


def send_email_otp(to_email: str, otp: str, purpose: str = "verification") -> bool:
    if not SMTP_USER or not SMTP_PASS or SMTP_USER == "your-email@gmail.com":
        print(f"[EMAIL OTP] SMTP not configured - would send to {to_email}: {otp}")
        return False
    try:
        msg = MIMEMultipart()
        msg["From"] = SMTP_FROM
        msg["To"] = to_email
        msg["Subject"] = f"Your verification code: {otp}"

        purpose_label = {
            "verification": "email verification",
            "reset_password": "password reset",
            "login": "login verification",
        }.get(purpose, purpose)

        body = f"""
        <html><body style="font-family: Arial, sans-serif; padding: 20px; background: #f5f5f5;">
            <div style="max-width: 500px; margin: 0 auto; background: white; border-radius: 12px;
                        padding: 30px; box-shadow: 0 4px 15px rgba(0,0,0,0.05);">
                <h2 style="color: #667eea; margin: 0 0 20px 0;">Student Marks Analyzer</h2>
                <p style="color: #333; font-size: 16px;">Your <b>{purpose_label}</b> code is:</p>
                <h1 style="background: linear-gradient(135deg, #667eea, #764ba2); color: white;
                           padding: 20px; border-radius: 12px; letter-spacing: 8px;
                           text-align: center; font-size: 32px; margin: 20px 0;">{otp}</h1>
                <p style="color: #666; font-size: 14px;">This code expires in <b>{OTP_EXPIRE_MINUTES} minutes</b>.</p>
            </div>
        </body></html>
        """
        msg.attach(MIMEText(body, "html"))
        with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=15) as server:
            server.starttls()
            server.login(SMTP_USER, SMTP_PASS)
            server.send_message(msg)
        print(f"[OK] Email OTP sent to {to_email}")
        return True
    except Exception as e:
        print(f"[ERROR] Email send failed: {type(e).__name__}: {e}")
        return False


def send_email_notification(to_email: str, subject: str, html_body: str) -> bool:
    if not SMTP_USER or not SMTP_PASS:
        print(f"[EMAIL] would send to {to_email}: {subject}")
        return False
    try:
        msg = MIMEMultipart()
        msg["From"] = SMTP_FROM
        msg["To"] = to_email
        msg["Subject"] = subject
        msg.attach(MIMEText(html_body, "html"))
        with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=15) as server:
            server.starttls()
            server.login(SMTP_USER, SMTP_PASS)
            server.send_message(msg)
        print(f"[OK] Email sent to {to_email}: {subject}")
        return True
    except Exception as e:
        print(f"[ERROR] Email failed: {e}")
        return False


def send_sms_otp(phone: str, otp: str) -> bool:
    if DEV_MODE_SMS or not FAST2SMS_API_KEY:
        print(f"[SMS OTP] {phone}: {otp} (dev mode)")
        return True
    try:
        import requests as req
        url = "https://www.fast2sms.com/dev/bulkV2"
        payload = {
            "route": "q",
            "message": f"Your OTP: {otp}. Valid {OTP_EXPIRE_MINUTES} min.",
            "language": "english", "flash": 0,
            "numbers": phone.lstrip("+91"),
        }
        headers = {"authorization": FAST2SMS_API_KEY, "Content-Type": "application/json"}
        r = req.post(url, json=payload, headers=headers, timeout=10)
        if r.status_code == 200:
            print(f"[OK] SMS OTP sent to {phone}")
            return True
        print(f"[ERROR] Fast2SMS: {r.text}")
        return False
    except Exception as e:
        print(f"[ERROR] SMS failed: {e}")
        return False


def send_whatsapp_otp(phone: str, otp: str) -> bool:
    if not TWILIO_AVAILABLE or not TWILIO_ACCOUNT_SID or not TWILIO_AUTH_TOKEN:
        print(f"[WHATSAPP] Not configured. Would send to {phone}: {otp}")
        return False
    try:
        client = TwilioClient(TWILIO_ACCOUNT_SID, TWILIO_AUTH_TOKEN)
        to_phone = phone if phone.startswith("+") else f"+{phone}"
        message = client.messages.create(
            body=f"Your Student Marks Analyzer OTP is: {otp}. Valid for {OTP_EXPIRE_MINUTES} minutes.",
            from_=TWILIO_WHATSAPP_FROM,
            to=f"whatsapp:{to_phone}",
        )
        print(f"[OK] WhatsApp OTP sent to {phone}: {message.sid}")
        return True
    except Exception as e:
        print(f"[ERROR] WhatsApp failed: {e}")
        return False


# ============================================================
# DATABASE
# ============================================================
DB_PATH = os.path.join(os.path.dirname(__file__), "student_marks.db")


@contextmanager
def get_db_connection():
    conn = sqlite3.connect(DB_PATH, timeout=30.0, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    try:
        conn.execute("PRAGMA journal_mode=WAL")
        conn.execute("PRAGMA busy_timeout=30000")
        conn.execute("PRAGMA synchronous=NORMAL")
    except Exception:
        pass
    try:
        yield conn
    finally:
        conn.close()


def init_database():
    with get_db_connection() as conn:
        cursor = conn.cursor()

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                email TEXT UNIQUE NOT NULL,
                phone TEXT,
                hashed_password TEXT NOT NULL,
                full_name TEXT,
                role TEXT DEFAULT 'teacher',
                is_active INTEGER DEFAULT 1,
                email_verified INTEGER DEFAULT 0,
                phone_verified INTEGER DEFAULT 0,
                twofa_enabled INTEGER DEFAULT 0,
                twofa_secret TEXT,
                language TEXT DEFAULT 'en',
                failed_login_attempts INTEGER DEFAULT 0,
                locked_until TIMESTAMP,
                last_login TIMESTAMP,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS otp_codes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                identifier TEXT NOT NULL,
                otp_code TEXT NOT NULL,
                purpose TEXT NOT NULL,
                expires_at TIMESTAMP NOT NULL,
                used INTEGER DEFAULT 0,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                name TEXT NOT NULL,
                marks TEXT NOT NULL,
                subjects TEXT NOT NULL,
                grade TEXT NOT NULL,
                average REAL NOT NULL,
                total_marks REAL NOT NULL,
                timestamp TEXT NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                semester TEXT,
                batch_year TEXT,
                department TEXT,
                class_name TEXT,
                photo_url TEXT,
                email TEXT,
                phone TEXT,
                date_of_birth TEXT,
                address TEXT
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS attendance (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER,
                date TEXT NOT NULL,
                status TEXT NOT NULL,
                subject TEXT,
                notes TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS performance_trends (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER,
                semester TEXT,
                average REAL,
                grade TEXT,
                total_marks REAL,
                timestamp TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS notifications (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER,
                user_id INTEGER,
                message TEXT,
                type TEXT,
                is_read INTEGER DEFAULT 0,
                email_sent INTEGER DEFAULT 0,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS parent_children (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                parent_user_id INTEGER NOT NULL,
                student_id INTEGER NOT NULL,
                relationship TEXT DEFAULT 'parent',
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(parent_user_id, student_id)
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS token_blacklist (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                token_hash TEXT UNIQUE NOT NULL,
                expires_at TIMESTAMP NOT NULL,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS login_attempts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT,
                ip_address TEXT,
                user_agent TEXT,
                success INTEGER,
                reason TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS audit_log (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                username TEXT,
                action TEXT NOT NULL,
                resource TEXT,
                details TEXT,
                ip_address TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS role_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                username TEXT NOT NULL,
                old_role TEXT NOT NULL,
                new_role TEXT NOT NULL,
                changed_by INTEGER,
                changed_by_username TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS classes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT UNIQUE NOT NULL,
                department TEXT,
                teacher_id INTEGER,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS email_queue (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                to_email TEXT NOT NULL,
                subject TEXT NOT NULL,
                html_body TEXT NOT NULL,
                sent INTEGER DEFAULT 0,
                attempts INTEGER DEFAULT 0,
                error TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS exams (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                class_name TEXT,
                subject TEXT,
                exam_date TEXT NOT NULL,
                start_time TEXT,
                end_time TEXT,
                total_marks INTEGER DEFAULT 100,
                room TEXT,
                notes TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS timetable (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                class_name TEXT NOT NULL,
                day_of_week TEXT NOT NULL,
                period INTEGER,
                subject TEXT NOT NULL,
                teacher_name TEXT,
                room TEXT,
                start_time TEXT,
                end_time TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS assignments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                description TEXT,
                class_name TEXT,
                subject TEXT,
                due_date TEXT NOT NULL,
                total_marks INTEGER DEFAULT 100,
                created_by INTEGER,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS assignment_submissions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                assignment_id INTEGER NOT NULL,
                student_id INTEGER NOT NULL,
                submitted_at TEXT,
                status TEXT DEFAULT 'pending',
                marks_obtained REAL,
                feedback TEXT,
                file_url TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(assignment_id, student_id)
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS fee_structures (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                class_name TEXT NOT NULL,
                fee_type TEXT NOT NULL,
                amount REAL NOT NULL,
                frequency TEXT DEFAULT 'monthly',
                academic_year TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS fee_payments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                fee_type TEXT NOT NULL,
                amount REAL NOT NULL,
                payment_date TEXT NOT NULL,
                payment_method TEXT DEFAULT 'cash',
                transaction_id TEXT,
                status TEXT DEFAULT 'paid',
                due_date TEXT,
                notes TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS saved_filters (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                name TEXT NOT NULL,
                entity TEXT NOT NULL,
                filter_json TEXT NOT NULL,
                is_shared INTEGER DEFAULT 0,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS backup_logs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                filename TEXT NOT NULL,
                size_bytes INTEGER,
                backup_type TEXT DEFAULT 'manual',
                created_by INTEGER,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS scheduled_reports (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                report_type TEXT NOT NULL,
                recipients TEXT NOT NULL,
                schedule TEXT NOT NULL,
                last_sent TIMESTAMP,
                enabled INTEGER DEFAULT 1,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        cursor.execute("""
            CREATE TABLE IF NOT EXISTS live_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                event_type TEXT NOT NULL,
                payload TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)

        # Default admin
        cursor.execute("SELECT COUNT(*) as count FROM users")
        if cursor.fetchone()["count"] == 0:
            default_hash = get_password_hash("Admin@123")
            cursor.execute("""
                INSERT INTO users (username, email, hashed_password, full_name, role,
                                   email_verified, phone_verified)
                VALUES (?, ?, ?, ?, ?, 1, 1)
            """, ("admin", "admin@example.com", default_hash, "Default Admin", "admin"))
            print("[OK] Default admin created - admin / Admin@123")

        cursor.execute("SELECT id FROM users WHERE username = 'admin2'")
        if not cursor.fetchone():
            backup_hash = get_password_hash("Admin2@123")
            cursor.execute("""
                INSERT INTO users (username, email, hashed_password, full_name, role,
                                   email_verified, phone_verified)
                VALUES (?, ?, ?, ?, ?, 1, 1)
            """, ("admin2", "admin2@example.com", backup_hash, "Backup Admin", "admin"))
            print("[OK] Backup admin created - admin2 / Admin2@123")

        cursor.execute("""
            UPDATE users SET role='admin', is_active=1, email_verified=1,
                             failed_login_attempts=0, locked_until=NULL
            WHERE username IN ('admin', 'admin2')
        """)

        for idx in [
            "CREATE INDEX IF NOT EXISTS idx_user_username ON users(username)",
            "CREATE INDEX IF NOT EXISTS idx_user_email ON users(email)",
            "CREATE INDEX IF NOT EXISTS idx_otp_identifier ON otp_codes(identifier)",
            "CREATE INDEX IF NOT EXISTS idx_student_class ON students(class_name)",
            "CREATE INDEX IF NOT EXISTS idx_notif_user ON notifications(user_id)",
            "CREATE INDEX IF NOT EXISTS idx_exams_date ON exams(exam_date)",
            "CREATE INDEX IF NOT EXISTS idx_assignments_due ON assignments(due_date)",
            "CREATE INDEX IF NOT EXISTS idx_fee_payments_student ON fee_payments(student_id)",
            "CREATE INDEX IF NOT EXISTS idx_live_events_user ON live_events(user_id)",
        ]:
            try:
                cursor.execute(idx)
            except sqlite3.OperationalError:
                pass

        conn.commit()
        print("[OK] Database initialized")


init_database()


# ============================================================
# MODELS (in main.py, existing)
# ============================================================
class UserRegister(BaseModel):
    username: str = Field(..., min_length=3, max_length=30)
    email: str
    phone: Optional[str] = None
    password: str = Field(..., min_length=8)
    full_name: Optional[str] = Field(None, max_length=100)
    role: str = "teacher"

    @validator("username")
    def v_username(cls, v):
        if not re.match(r"^[a-zA-Z0-9_]+$", v):
            raise ValueError("Username: letters, numbers, underscores only")
        return v.lower()

    @validator("role")
    def v_role(cls, v):
        if v not in ["admin", "teacher", "student", "parent"]:
            raise ValueError("Invalid role")
        return v


class OTPRequest(BaseModel):
    identifier: str
    purpose: str = "verification"


class OTPVerify(BaseModel):
    identifier: str
    otp_code: str
    purpose: str = "verification"


class ResendOTPRequest(BaseModel):
    email: str


class PasswordResetOTP(BaseModel):
    identifier: str
    otp_code: str
    new_password: str


class PasswordChange(BaseModel):
    old_password: str
    new_password: str


class MarksRequest(BaseModel):
    marks: List[float]
    student_name: Optional[str] = None
    subject_names: Optional[List[str]] = None
    passing_threshold: float = 40
    grade_scheme: str = "standard"
    semester: Optional[str] = None
    batch_year: Optional[str] = None
    department: Optional[str] = None
    class_name: Optional[str] = None

    @validator("marks")
    def v_marks(cls, v):
        if not v:
            raise ValueError("Marks cannot be empty")
        if any(m < 0 or m > 100 for m in v):
            raise ValueError("Marks must be 0-100")
        return v

    @validator("subject_names")
    def v_subjects(cls, v):
        if v:
            for name in v:
                ok, err = validate_subject_name(name)
                if not ok:
                    raise ValueError(f"Invalid subject '{name}': {err}")
        return v


class StudentRecord(BaseModel):
    name: str
    marks: List[float]
    subjects: List[str]
    grade: str
    average: float
    total_marks: Optional[float] = None
    timestamp: str
    semester: Optional[str] = None
    batch_year: Optional[str] = None
    department: Optional[str] = None
    class_name: Optional[str] = None

    @validator("subjects")
    def v_sub(cls, v):
        for name in v:
            ok, err = validate_subject_name(name)
            if not ok:
                raise ValueError(f"Invalid subject '{name}': {err}")
        return v


class AttendanceRecord(BaseModel):
    student_id: int
    date: str
    status: str
    subject: Optional[str] = None
    notes: Optional[str] = None


class NotificationCreate(BaseModel):
    student_id: Optional[int] = None
    user_id: Optional[int] = None
    message: str
    type: str = "Info"


class CompareRequest(BaseModel):
    student_ids: List[int]


class LinkChildRequest(BaseModel):
    parent_user_id: int
    student_id: int
    relationship: Optional[str] = "parent"


class RoleChangeRequest(BaseModel):
    role: str


class BulkDeleteRequest(BaseModel):
    student_ids: List[int]


class BulkNotificationRequest(BaseModel):
    student_ids: List[int]
    message: str
    type: str = "Info"


class BulkAttendanceRequest(BaseModel):
    student_ids: List[int]
    date: str
    status: str
    subject: Optional[str] = None


class ClassCreate(BaseModel):
    name: str
    department: Optional[str] = None
    teacher_id: Optional[int] = None


class ClassAssignRequest(BaseModel):
    student_ids: List[int]
    class_name: str


class LanguageUpdate(BaseModel):
    language: str


class TOTPVerify(BaseModel):
    code: str


# ============================================================
# HELPERS
# ============================================================
def hash_token(token: str) -> str:
    import hashlib
    return hashlib.sha256(token.encode()).hexdigest()


def blacklist_token(token: str, expires_at: datetime):
    try:
        with get_db_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(
                "INSERT OR IGNORE INTO token_blacklist (token_hash, expires_at) VALUES (?, ?)",
                (hash_token(token), expires_at.isoformat()),
            )
            conn.commit()
    except Exception:
        pass


def is_token_blacklisted(token: str) -> bool:
    try:
        with get_db_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT id FROM token_blacklist WHERE token_hash=?", (hash_token(token),))
            return cursor.fetchone() is not None
    except Exception:
        return False


def log_audit(user_id, username, action, resource=None, details=None, ip=None, conn=None):
    if conn is not None:
        try:
            cursor = conn.cursor()
            cursor.execute(
                """INSERT INTO audit_log (user_id, username, action, resource, details, ip_address)
                   VALUES (?, ?, ?, ?, ?, ?)""",
                (user_id, username, action, resource, details, ip),
            )
            conn.commit()
        except Exception:
            pass
        return
    try:
        with get_db_connection() as c:
            cursor = c.cursor()
            cursor.execute(
                """INSERT INTO audit_log (user_id, username, action, resource, details, ip_address)
                   VALUES (?, ?, ?, ?, ?, ?)""",
                (user_id, username, action, resource, details, ip),
            )
            c.commit()
    except Exception:
        pass


def log_login_attempt(username, ip, ua, success, reason="", conn=None):
    if conn is not None:
        try:
            cursor = conn.cursor()
            cursor.execute(
                """INSERT INTO login_attempts (username, ip_address, user_agent, success, reason)
                   VALUES (?, ?, ?, ?, ?)""",
                (username, ip, (ua or "")[:200], 1 if success else 0, reason),
            )
            conn.commit()
        except Exception:
            pass
        return
    try:
        with get_db_connection() as c:
            cursor = c.cursor()
            cursor.execute(
                """INSERT INTO login_attempts (username, ip_address, user_agent, success, reason)
                   VALUES (?, ?, ?, ?, ?)""",
                (username, ip, (ua or "")[:200], 1 if success else 0, reason),
            )
            c.commit()
    except Exception:
        pass


def check_account_locked(username, conn=None):
    try:
        if conn is not None:
            cursor = conn.cursor()
            cursor.execute("SELECT locked_until FROM users WHERE username=?", (username,))
            row = cursor.fetchone()
        else:
            with get_db_connection() as c:
                cursor = c.cursor()
                cursor.execute("SELECT locked_until FROM users WHERE username=?", (username,))
                row = cursor.fetchone()
        if row and row["locked_until"]:
            try:
                lu = datetime.fromisoformat(row["locked_until"])
                if datetime.utcnow() < lu:
                    mins = int((lu - datetime.utcnow()).total_seconds() / 60) + 1
                    return f"Account locked. Try again in {mins} minute(s)."
            except ValueError:
                pass
    except Exception:
        pass
    return None


def increment_failed_attempts(username, conn=None):
    if conn is not None:
        _do_increment(conn, username)
        return
    with get_db_connection() as c:
        _do_increment(c, username)


def _do_increment(conn, username):
    try:
        cursor = conn.cursor()
        cursor.execute("SELECT failed_login_attempts FROM users WHERE username=?", (username,))
        row = cursor.fetchone()
        if not row:
            return
        attempts = (row["failed_login_attempts"] or 0) + 1
        if attempts >= MAX_FAILED_ATTEMPTS:
            lu = (datetime.utcnow() + timedelta(minutes=LOCKOUT_DURATION_MINUTES)).isoformat()
            cursor.execute(
                "UPDATE users SET failed_login_attempts=?, locked_until=? WHERE username=?",
                (attempts, lu, username),
            )
        else:
            cursor.execute(
                "UPDATE users SET failed_login_attempts=? WHERE username=?",
                (attempts, username),
            )
        conn.commit()
    except Exception:
        pass


def reset_failed_attempts(username, conn=None):
    if conn is not None:
        try:
            cursor = conn.cursor()
            cursor.execute(
                """UPDATE users SET failed_login_attempts=0, locked_until=NULL, last_login=?
                   WHERE username=?""",
                (datetime.utcnow().isoformat(), username),
            )
            conn.commit()
        except Exception:
            pass
        return
    with get_db_connection() as c:
        cursor = c.cursor()
        cursor.execute(
            """UPDATE users SET failed_login_attempts=0, locked_until=NULL, last_login=?
               WHERE username=?""",
            (datetime.utcnow().isoformat(), username),
        )
        c.commit()


def save_otp(identifier, otp, purpose):
    expires = (datetime.utcnow() + timedelta(minutes=OTP_EXPIRE_MINUTES)).isoformat()
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "UPDATE otp_codes SET used=1 WHERE identifier=? AND purpose=? AND used=0",
            (identifier, purpose),
        )
        cursor.execute(
            """INSERT INTO otp_codes (identifier, otp_code, purpose, expires_at)
               VALUES (?, ?, ?, ?)""",
            (identifier, otp, purpose, expires),
        )
        conn.commit()


def verify_otp(identifier, otp, purpose):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            """SELECT * FROM otp_codes
               WHERE identifier=? AND purpose=? AND used=0
               ORDER BY created_at DESC LIMIT 1""",
            (identifier, purpose),
        )
        row = cursor.fetchone()
        if not row:
            return False, "No OTP found. Request a new one."
        if row["otp_code"] != otp:
            return False, "Incorrect OTP"
        try:
            exp = datetime.fromisoformat(row["expires_at"])
            if datetime.utcnow() > exp:
                return False, "OTP expired. Request a new one."
        except ValueError:
            pass
        cursor.execute("UPDATE otp_codes SET used=1 WHERE id=?", (row["id"],))
        conn.commit()
        return True, "OK"


def queue_email(to_email: str, subject: str, html_body: str):
    try:
        with get_db_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(
                "INSERT INTO email_queue (to_email, subject, html_body) VALUES (?, ?, ?)",
                (to_email, subject, html_body),
            )
            conn.commit()
    except Exception as e:
        print(f"Queue error: {e}")


def send_welcome_email(user_email: str, username: str, role: str):
    subject = "Welcome to Student Marks Analyzer Pro"
    html = f"""
    <html><body style="font-family: Arial, sans-serif; padding: 20px;">
        <h2 style="color: #667eea;">Welcome, {username}!</h2>
        <p>Your account has been created with role: <b>{role}</b>.</p>
        <p>You can now log in and start using the app.</p>
    </body></html>
    """
    queue_email(user_email, subject, html)


def send_grade_notification_email(student_name: str, grade: str, average: float, parent_emails: List[str]):
    subject = f"Grade Update: {student_name} - {grade}"
    html = f"""
    <html><body style="font-family: Arial, sans-serif; padding: 20px;">
        <h2 style="color: #667eea;">Grade Update</h2>
        <p><b>Student:</b> {student_name}</p>
        <p><b>Grade:</b> {grade}</p>
        <p><b>Average:</b> {average:.2f}%</p>
    </body></html>
    """
    for email in parent_emails:
        if email:
            queue_email(email, subject, html)


def process_email_queue():
    try:
        with get_db_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                SELECT id, to_email, subject, html_body, attempts
                FROM email_queue WHERE sent=0 AND attempts < 3
                ORDER BY created_at ASC LIMIT 10
            """)
            emails = cursor.fetchall()
            for email in emails:
                success = send_email_notification(
                    email["to_email"], email["subject"], email["html_body"]
                )
                if success:
                    cursor.execute("UPDATE email_queue SET sent=1 WHERE id=?", (email["id"],))
                else:
                    cursor.execute(
                        "UPDATE email_queue SET attempts=attempts+1 WHERE id=?",
                        (email["id"],)
                    )
            conn.commit()
    except Exception as e:
        print(f"Queue processing error: {e}")


# ============================================================
# ANALYTICS HELPERS
# ============================================================
def calculate_statistics(marks):
    arr = np.array(marks)
    return {
        "total": float(np.sum(arr)),
        "average": float(np.mean(arr)),
        "highest": float(np.max(arr)),
        "lowest": float(np.min(arr)),
        "median": float(np.median(arr)),
        "std_deviation": float(np.std(arr)),
        "range": float(np.max(arr) - np.min(arr)),
        "variance": float(np.var(arr)),
        "quartiles": {
            "q1": float(np.percentile(arr, 25)),
            "q2": float(np.percentile(arr, 50)),
            "q3": float(np.percentile(arr, 75)),
        },
    }


def determine_grade(avg, scheme="standard"):
    schemes = {
        "standard": {(90, 100): ("A+", 10), (80, 89.99): ("A", 9), (70, 79.99): ("B", 8),
                     (60, 69.99): ("C", 7), (50, 59.99): ("D", 6), (40, 49.99): ("E", 5),
                     (0, 39.99): ("F", 0)},
        "strict": {(93, 100): ("A+", 10), (85, 92.99): ("A", 9), (75, 84.99): ("B", 8),
                   (65, 74.99): ("C", 7), (55, 64.99): ("D", 6), (45, 54.99): ("E", 5),
                   (0, 44.99): ("F", 0)},
        "lenient": {(85, 100): ("A+", 10), (70, 84.99): ("A", 9), (60, 69.99): ("B", 8),
                    (50, 59.99): ("C", 7), (40, 49.99): ("D", 6), (30, 39.99): ("E", 5),
                    (0, 29.99): ("F", 0)},
    }
    for (lo, hi), (g, p) in schemes.get(scheme, schemes["standard"]).items():
        if lo <= avg <= hi:
            return g, p
    return "F", 0


def generate_ai_insights(marks, avg, subjects):
    ins = {"strengths": [], "weaknesses": [], "recommendations": [],
           "predictions": [], "study_tips": [], "improvement_potential": 0}
    for i, m in enumerate(marks):
        s = subjects[i] if i < len(subjects) else f"Subject {i+1}"
        if m >= 80:
            ins["strengths"].append({"subject": s, "marks": m, "level": "Excellent"})
        elif m < 40:
            ins["weaknesses"].append({"subject": s, "marks": m, "level": "Needs Improvement"})
    if avg < 40:
        ins["recommendations"].append({"priority": "High", "action": "Immediate intervention",
                                       "details": "Consider tutoring"})
        ins["study_tips"] = ["Focus on fundamentals", "Practice daily"]
    elif avg < 60:
        ins["recommendations"].append({"priority": "Medium", "action": "Structured plan",
                                       "details": "Create schedule"})
        ins["study_tips"] = ["Join study groups", "Use online resources"]
    elif avg < 75:
        ins["recommendations"].append({"priority": "Low", "action": "Enhance performance",
                                       "details": "Focus on weak areas"})
        ins["study_tips"] = ["Advanced problems", "Teach peers"]
    else:
        ins["recommendations"].append({"priority": "Optional", "action": "Excellence program",
                                       "details": "Advanced topics"})
        ins["study_tips"] = ["Peer tutoring", "Competitions"]
    if len(marks) >= 5:
        trend = np.polyfit(range(len(marks)), marks, 1)[0]
        ins["predictions"].append("Positive trend" if trend > 0
                                   else "Declining trend" if trend < 0 else "Stable")
    max_total = len(marks) * 100
    ins["improvement_potential"] = round(((max_total - sum(marks)) / max_total) * 100, 2)
    return ins


def detect_anomalies(marks):
    anomalies = []
    arr = np.array(marks)
    if len(arr) < 3:
        return anomalies
    mean, std = np.mean(arr), np.std(arr)
    for i, m in enumerate(marks):
        z = (m - mean) / std if std > 0 else 0
        if abs(z) > 2:
            anomalies.append({"subject": f"Subject {i+1}", "marks": m, "z_score": round(z, 2),
                              "type": "Exceptionally High" if m > mean else "Exceptionally Low",
                              "suggestion": "Review"})
    return anomalies


def generate_recommendations(marks, avg):
    recs = []
    if avg < 40:
        recs += ["Urgent: Consider additional tutoring", "Focus on foundational concepts"]
    elif avg < 60:
        recs += ["Need improvement: Study groups recommended", "Create structured schedule"]
    elif avg < 75:
        recs += ["Good performance: Focus on weak areas", "Set higher targets"]
    else:
        recs += ["Excellent! Help peers", "Aim for top performance"]
    weak = [i for i, m in enumerate(marks) if m < 40]
    if weak:
        recs.append(f"Focus on subjects {', '.join(str(i+1) for i in weak)}")
    strong = [i for i, m in enumerate(marks) if m >= 80]
    if strong:
        recs.append(f"Strong in subjects {', '.join(str(i+1) for i in strong)}")
    return recs


def analyze_trends(student_id):
    try:
        with get_db_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(
                "SELECT * FROM performance_trends WHERE student_id=? ORDER BY created_at ASC",
                (student_id,),
            )
            trends = [dict(r) for r in cursor.fetchall()]
            if len(trends) < 2:
                return {"message": "Not enough data", "trends": trends}
            avgs = [t["average"] for t in trends]
            imp = avgs[-1] - avgs[0]
            d = "Improving" if imp > 5 else "Declining" if imp < -5 else "Stable"
            return {"student_id": student_id, "total_records": len(trends),
                    "first_average": avgs[0], "current_average": avgs[-1],
                    "improvement": round(imp, 2), "trend_direction": d, "trends": trends}
    except Exception as e:
        return {"error": str(e)}


def get_parent_emails_for_student(student_id: int) -> List[str]:
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT u.email FROM users u
            JOIN parent_children pc ON u.id = pc.parent_user_id
            WHERE pc.student_id = ? AND u.is_active = 1
        """, (student_id,))
        return [r["email"] for r in cursor.fetchall()]


# ============================================================
# WEBSOCKET MANAGER
# ============================================================
class ConnectionManager:
    def __init__(self):
        self.active_connections: Dict[int, List[WebSocket]] = {}

    async def connect(self, websocket: WebSocket, user_id: int):
        await websocket.accept()
        self.active_connections.setdefault(user_id, []).append(websocket)
        print(f"[WS] User {user_id} connected")

    def disconnect(self, websocket: WebSocket, user_id: int):
        if user_id in self.active_connections:
            try:
                self.active_connections[user_id].remove(websocket)
            except ValueError:
                pass
            if not self.active_connections[user_id]:
                del self.active_connections[user_id]

    async def send_to_user(self, user_id: int, message: dict):
        if user_id in self.active_connections:
            for connection in list(self.active_connections[user_id]):
                try:
                    await connection.send_json(message)
                except Exception:
                    self.disconnect(connection, user_id)

    async def broadcast(self, message: dict):
        for user_id in list(self.active_connections.keys()):
            await self.send_to_user(user_id, message)


manager = ConnectionManager()


def log_live_event(user_id: int, event_type: str, payload: dict):
    try:
        with get_db_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                INSERT INTO live_events (user_id, event_type, payload)
                VALUES (?, ?, ?)
            """, (user_id, event_type, json.dumps(payload)))
            conn.commit()
    except Exception:
        pass


# ============================================================
# AUTH DEPENDENCIES
# ============================================================
async def get_current_user(token: str = Depends(oauth2_scheme)):
    if not token:
        return None
    if is_token_blacklisted(token):
        return None
    payload = decode_token(token)
    if not payload or payload.get("type") != "access":
        return None
    username = payload.get("sub")
    if not username:
        return None
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM users WHERE username=? AND is_active=1", (username,))
        row = cursor.fetchone()
        return dict(row) if row else None


async def require_user(user=Depends(get_current_user)):
    if user is None:
        raise HTTPException(401, "Authentication required",
                            headers={"WWW-Authenticate": "Bearer"})
    return user


async def require_admin(user=Depends(require_user)):
    if user.get("role") != "admin":
        raise HTTPException(403, "Admin access required")
    return user


def require_role(*roles):
    async def checker(user=Depends(require_user)):
        if user.get("role") not in roles:
            raise HTTPException(403, f"Requires role: {', '.join(roles)}")
        return user
    return checker
# ============================================================
# AUTH ENDPOINTS
# ============================================================
@app.post("/auth/register")
@limiter.limit("10/minute")
def register_user(request: Request, user: UserRegister, background_tasks: BackgroundTasks):
    is_valid, error = validate_password_strength(user.password)
    if not is_valid:
        raise HTTPException(400, error)

    if not validate_email(user.email):
        raise HTTPException(400, "Invalid email format")

    if not SMTP_USER or not SMTP_PASS or SMTP_USER == "your-email@gmail.com":
        raise HTTPException(500, "Email service not configured.")

    phone_clean = None
    if user.phone:
        ok, result = validate_phone(user.phone)
        if not ok:
            raise HTTPException(400, result)
        phone_clean = result

    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "SELECT id FROM users WHERE username=? OR email=?",
            (user.username, user.email.lower()),
        )
        if cursor.fetchone():
            raise HTTPException(400, "Username or email already exists")

        hashed = get_password_hash(user.password)
        cursor.execute(
            """INSERT INTO users (username, email, phone, hashed_password, full_name,
                                  role, email_verified, phone_verified)
               VALUES (?, ?, ?, ?, ?, ?, 0, 0)""",
            (user.username, user.email.lower(), phone_clean, hashed,
             user.full_name or user.username, user.role),
        )
        user_id = cursor.lastrowid
        conn.commit()

    otp = generate_otp()
    save_otp(user.email.lower(), otp, "verification")
    sent = send_email_otp(user.email, otp, "verification")

    if not sent:
        with get_db_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("DELETE FROM users WHERE id=?", (user_id,))
            conn.commit()
        raise HTTPException(500, "Could not send verification email.")

    if phone_clean:
        phone_otp = generate_otp()
        save_otp(phone_clean, phone_otp, "phone_verification")
        send_sms_otp(phone_clean, phone_otp)

    background_tasks.add_task(send_welcome_email, user.email.lower(), user.username, user.role)

    log_audit(user_id, user.username, "register", "user",
              f"Role: {user.role}",
              request.client.host if request.client else None)

    return {
        "message": "Registration successful. Check your email for the verification code.",
        "user_id": user_id,
        "username": user.username,
        "email": user.email.lower(),
        "phone": phone_clean,
        "email_verification_required": True,
        "phone_verification_required": bool(phone_clean),
    }


@app.post("/auth/resend-otp")
@limiter.limit("3/minute")
def resend_otp(request: Request, data: ResendOTPRequest):
    email = data.email.lower().strip()
    if not validate_email(email):
        raise HTTPException(400, "Invalid email")

    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT id, email_verified FROM users WHERE email=?", (email,))
        user = cursor.fetchone()

    if not user:
        raise HTTPException(404, "No account found")
    if user["email_verified"]:
        raise HTTPException(400, "Already verified")

    otp = generate_otp()
    save_otp(email, otp, "verification")
    sent = send_email_otp(email, otp, "verification")

    if not sent:
        raise HTTPException(500, "Could not send OTP")

    return {"message": "OTP sent to your email", "delivery": "email"}


@app.post("/auth/verify-otp")
def verify_otp_endpoint(data: OTPVerify):
    ok, msg = verify_otp(data.identifier.strip().lower(), data.otp_code, data.purpose)
    if not ok:
        raise HTTPException(400, msg)

    with get_db_connection() as conn:
        cursor = conn.cursor()
        if data.purpose == "verification":
            cursor.execute(
                "UPDATE users SET email_verified=1 WHERE email=?",
                (data.identifier.strip().lower(),),
            )
        elif data.purpose == "phone_verification":
            ok_p, phone = validate_phone(data.identifier)
            if ok_p:
                cursor.execute("UPDATE users SET phone_verified=1 WHERE phone=?", (phone,))
        conn.commit()

    return {"message": "Verification successful", "verified": True}


@app.post("/auth/login")
@limiter.limit("10/minute")
def login(request: Request, username: str = "", password: str = "",
          email: str = "", pwd: str = "", totp_code: str = ""):
    ip = request.client.host if request.client else "unknown"
    ua = request.headers.get("user-agent", "")

    identifier = (username or email or "").strip().lower()
    pwd_value = password or pwd or ""

    if not identifier or not pwd_value:
        raise HTTPException(400, "Username/email and password required")

    with get_db_connection() as conn:
        lock_msg = check_account_locked(identifier, conn)
        if lock_msg:
            log_login_attempt(identifier, ip, ua, False, "account_locked", conn)
            raise HTTPException(423, lock_msg)

        cursor = conn.cursor()
        cursor.execute(
            "SELECT * FROM users WHERE username=? OR email=?",
            (identifier, identifier),
        )
        user = cursor.fetchone()

        if not user:
            log_login_attempt(identifier, ip, ua, False, "user_not_found", conn)
            raise HTTPException(401, "Invalid credentials")

        user = dict(user)

        if not user.get("is_active", 1):
            log_login_attempt(identifier, ip, ua, False, "account_disabled", conn)
            raise HTTPException(403, "Account disabled")

        if not verify_password(pwd_value, user["hashed_password"]):
            increment_failed_attempts(user["username"], conn)
            log_login_attempt(identifier, ip, ua, False, "wrong_password", conn)
            raise HTTPException(401, "Invalid credentials")

        if not user.get("email_verified") and user["username"] not in ("admin", "admin2"):
            log_login_attempt(identifier, ip, ua, False, "email_not_verified", conn)
            raise HTTPException(403, "Email not verified. Please check your inbox.")

        if user.get("twofa_enabled") and user.get("twofa_secret"):
            if not totp_code:
                return {"requires_2fa": True, "message": "Enter your 6-digit 2FA code"}
            totp = pyotp.TOTP(user["twofa_secret"])
            if not totp.verify(totp_code, valid_window=1):
                log_login_attempt(identifier, ip, ua, False, "invalid_2fa", conn)
                raise HTTPException(401, "Invalid 2FA code")

        reset_failed_attempts(user["username"], conn)
        log_login_attempt(identifier, ip, ua, True, "success", conn)
        log_audit(user["id"], user["username"], "login", "user", None, ip, conn)

    access_token = create_access_token({
        "sub": user["username"], "role": user["role"], "user_id": user["id"],
    })
    refresh = create_refresh_token({"sub": user["username"], "user_id": user["id"]})

    return {
        "access_token": access_token,
        "refresh_token": refresh,
        "token_type": "bearer",
        "expires_in": ACCESS_TOKEN_EXPIRE_MINUTES * 60,
        "user": {
            "id": user["id"],
            "username": user["username"],
            "email": user["email"],
            "phone": user.get("phone"),
            "full_name": user.get("full_name"),
            "role": user["role"],
            "language": user.get("language", "en"),
            "email_verified": bool(user.get("email_verified")),
            "phone_verified": bool(user.get("phone_verified")),
            "twofa_enabled": bool(user.get("twofa_enabled")),
        },
    }


@app.post("/auth/send-otp")
@limiter.limit("5/minute")
def send_otp(request: Request, data: OTPRequest):
    identifier = data.identifier.strip().lower()

    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "SELECT * FROM users WHERE email=? OR phone=?",
            (identifier, data.identifier.strip()),
        )
        user = cursor.fetchone()

    if not user and data.purpose == "reset_password":
        raise HTTPException(404, "No account found")

    otp = generate_otp()

    if "@" in identifier:
        save_otp(identifier, otp, data.purpose)
        sent = send_email_otp(identifier, otp, data.purpose)
        if not sent:
            raise HTTPException(500, "Could not send OTP")
        return {"message": "OTP sent to your email", "delivery": "email"}
    else:
        ok, phone = validate_phone(identifier)
        if not ok:
            raise HTTPException(400, phone)
        save_otp(phone, otp, data.purpose)
        sent = send_sms_otp(phone, otp)
        return {"message": "OTP sent", "delivery": "sms" if sent else "console"}


@app.post("/auth/whatsapp-otp")
@limiter.limit("5/minute")
def send_whatsapp_otp_endpoint(request: Request, data: WhatsAppOTPRequest):
    ok, phone = validate_phone(data.phone)
    if not ok:
        raise HTTPException(400, phone)

    otp = generate_otp()
    save_otp(phone, otp, data.purpose)
    sent = send_whatsapp_otp(phone, otp)

    if not sent:
        raise HTTPException(500, "WhatsApp sending failed. Check Twilio configuration.")

    return {
        "message": "OTP sent via WhatsApp",
        "delivery": "whatsapp",
        "phone": phone,
    }


@app.post("/auth/reset-password")
@limiter.limit("5/minute")
def reset_password(request: Request, data: PasswordResetOTP):
    ok, msg = verify_otp(data.identifier.strip().lower(), data.otp_code, "reset_password")
    if not ok:
        raise HTTPException(400, msg)

    is_valid, error = validate_password_strength(data.new_password)
    if not is_valid:
        raise HTTPException(400, error)

    hashed = get_password_hash(data.new_password)
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            """UPDATE users SET hashed_password=?, failed_login_attempts=0, locked_until=NULL
               WHERE email=? OR phone=?""",
            (hashed, data.identifier.strip().lower(), data.identifier.strip()),
        )
        if cursor.rowcount == 0:
            raise HTTPException(404, "User not found")
        conn.commit()

    return {"message": "Password reset successfully. Please log in."}


@app.get("/auth/me")
def get_me(user=Depends(require_user)):
    return {
        "id": user["id"], "username": user["username"], "email": user["email"],
        "phone": user.get("phone"), "full_name": user.get("full_name"),
        "role": user["role"],
        "language": user.get("language", "en"),
        "email_verified": bool(user.get("email_verified")),
        "phone_verified": bool(user.get("phone_verified")),
        "twofa_enabled": bool(user.get("twofa_enabled")),
        "last_login": user.get("last_login"),
        "created_at": user.get("created_at"),
    }


@app.post("/auth/change-password")
def change_password(data: PasswordChange, user=Depends(require_user)):
    if not verify_password(data.old_password, user["hashed_password"]):
        raise HTTPException(400, "Current password is incorrect")
    is_valid, error = validate_password_strength(data.new_password)
    if not is_valid:
        raise HTTPException(400, error)
    new_hash = get_password_hash(data.new_password)
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("UPDATE users SET hashed_password=? WHERE id=?", (new_hash, user["id"]))
        conn.commit()
    return {"message": "Password changed successfully"}


@app.post("/auth/logout")
def logout(token: str = Depends(oauth2_scheme), user=Depends(require_user)):
    if token:
        payload = decode_token(token)
        if payload and payload.get("exp"):
            blacklist_token(token, datetime.utcfromtimestamp(payload["exp"]))
    return {"message": "Logged out"}


@app.post("/auth/check-strength")
def check_strength(password: str = Query(..., min_length=1)):
    score = 0
    if len(password) >= 8: score += 20
    if len(password) >= 12: score += 10
    if re.search(r"[A-Z]", password): score += 15
    if re.search(r"[a-z]", password): score += 15
    if re.search(r"\d", password): score += 15
    if re.search(r"[!@#$%^&*(),.?\":{}|<>_\-+=\[\]\\\/;'`~]", password): score += 15
    if len(set(password)) > 6: score += 10
    score = min(score, 100)
    label = "Weak" if score < 40 else "Fair" if score < 60 else "Good" if score < 80 else "Strong"
    valid, err = validate_password_strength(password)
    return {"score": score, "label": label, "valid": valid, "error": err}


# ============================================================
# 2FA
# ============================================================
@app.post("/auth/2fa/setup")
def setup_2fa(user=Depends(require_user)):
    secret = pyotp.random_base32()
    totp_uri = pyotp.totp.TOTP(secret).provisioning_uri(
        name=user["email"], issuer_name=TOTP_ISSUER
    )

    img = qrcode.make(totp_uri)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    import base64
    qr_b64 = base64.b64encode(buf.getvalue()).decode()

    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("UPDATE users SET twofa_secret=? WHERE id=?", (secret, user["id"]))
        conn.commit()

    return {
        "secret": secret,
        "qr_code": f"data:image/png;base64,{qr_b64}",
        "uri": totp_uri,
    }


@app.post("/auth/2fa/verify")
def verify_2fa(data: TOTPVerify, user=Depends(require_user)):
    if not user.get("twofa_secret"):
        raise HTTPException(400, "Call /auth/2fa/setup first")

    totp = pyotp.TOTP(user["twofa_secret"])
    if not totp.verify(data.code, valid_window=1):
        raise HTTPException(400, "Invalid code")

    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("UPDATE users SET twofa_enabled=1 WHERE id=?", (user["id"],))
        conn.commit()

    return {"message": "2FA enabled successfully", "enabled": True}


@app.post("/auth/2fa/disable")
def disable_2fa(data: TOTPVerify, user=Depends(require_user)):
    if not user.get("twofa_enabled"):
        raise HTTPException(400, "2FA not enabled")

    totp = pyotp.TOTP(user["twofa_secret"])
    if not totp.verify(data.code, valid_window=1):
        raise HTTPException(400, "Invalid code")

    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "UPDATE users SET twofa_enabled=0, twofa_secret=NULL WHERE id=?",
            (user["id"],)
        )
        conn.commit()

    return {"message": "2FA disabled", "enabled": False}


# ============================================================
# USER MANAGEMENT
# ============================================================
@app.get("/auth/users")
def list_users(admin=Depends(require_admin)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT id, username, email, phone, full_name, role, is_active,
                   email_verified, phone_verified, twofa_enabled, language,
                   last_login, created_at
            FROM users ORDER BY id
        """)
        return {"users": [dict(r) for r in cursor.fetchall()]}


@app.put("/auth/users/{user_id}/toggle")
def toggle_user(request: Request, user_id: int, admin=Depends(require_admin)):
    if user_id == admin["id"]:
        raise HTTPException(400, "You cannot disable your own account.")

    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT is_active, username FROM users WHERE id=?", (user_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(404, "User not found")

        new_status = 0 if row["is_active"] else 1
        cursor.execute("UPDATE users SET is_active=? WHERE id=?", (new_status, user_id))
        conn.commit()

    log_audit(admin["id"], admin["username"], "toggle_user", f"user:{user_id}",
              f"new_status={new_status}",
              request.client.host if request.client else None)
    return {"message": "Toggled", "user_id": user_id, "is_active": new_status}


@app.put("/auth/users/{user_id}/role")
def change_role(request: Request, user_id: int, data: RoleChangeRequest,
                admin=Depends(require_admin)):
    role = data.role
    if role not in ["admin", "teacher", "student", "parent"]:
        raise HTTPException(400, "Invalid role")

    if user_id == admin["id"] and role != "admin":
        raise HTTPException(400, "You cannot demote yourself. Ask another admin.")

    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT username, role FROM users WHERE id=?", (user_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(404, "User not found")

        old_role = row["role"]
        username = row["username"]
        if old_role == role:
            raise HTTPException(400, f"Already has role '{role}'")

        cursor.execute("UPDATE users SET role=? WHERE id=?", (role, user_id))
        cursor.execute("""
            INSERT INTO role_history (user_id, username, old_role, new_role,
                                      changed_by, changed_by_username)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (user_id, username, old_role, role, admin["id"], admin["username"]))
        conn.commit()

    log_audit(admin["id"], admin["username"], "change_role", f"user:{user_id}",
              f"{old_role} -> {role}",
              request.client.host if request.client else None)

    return {"message": f"Role changed: {old_role} -> {role}",
            "user_id": user_id, "old_role": old_role, "new_role": role}


@app.get("/auth/users/{user_id}/role-history")
def get_role_history(user_id: int, admin=Depends(require_admin)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "SELECT * FROM role_history WHERE user_id=? ORDER BY created_at DESC LIMIT 50",
            (user_id,),
        )
        return {"history": [dict(r) for r in cursor.fetchall()]}


@app.put("/auth/users/{user_id}/revert-role")
def revert_role(request: Request, user_id: int, admin=Depends(require_admin)):
    if user_id == admin["id"]:
        raise HTTPException(400, "You cannot revert your own role.")

    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "SELECT old_role FROM role_history WHERE user_id=? ORDER BY created_at DESC LIMIT 1",
            (user_id,),
        )
        last = cursor.fetchone()
        if not last:
            raise HTTPException(400, "No previous role change")

        cursor.execute("SELECT username, role FROM users WHERE id=?", (user_id,))
        user = cursor.fetchone()
        if not user:
            raise HTTPException(404, "User not found")

        current = user["role"]
        prev = last["old_role"]
        if current == prev:
            raise HTTPException(400, f"Already at '{prev}'")

        cursor.execute("UPDATE users SET role=? WHERE id=?", (prev, user_id))
        cursor.execute("""
            INSERT INTO role_history (user_id, username, old_role, new_role,
                                      changed_by, changed_by_username)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (user_id, user["username"], current, prev, admin["id"], admin["username"]))
        conn.commit()

    return {"message": f"Reverted: {current} -> {prev}",
            "user_id": user_id, "old_role": current, "new_role": prev}


@app.delete("/auth/users/{user_id}")
def delete_user(request: Request, user_id: int, admin=Depends(require_admin)):
    if user_id == admin["id"]:
        raise HTTPException(400, "Cannot delete your own account")
    if user_id in (1, 2):
        raise HTTPException(400, "Cannot delete default admin accounts")
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT username FROM users WHERE id=?", (user_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(404, "User not found")
        cursor.execute("DELETE FROM parent_children WHERE parent_user_id=?", (user_id,))
        cursor.execute("DELETE FROM users WHERE id=?", (user_id,))
        conn.commit()
    log_audit(admin["id"], admin["username"], "delete_user", f"user:{user_id}",
              f"deleted:{row['username']}",
              request.client.host if request.client else None)
    return {"message": "User deleted"}


@app.put("/auth/language")
def update_language(data: LanguageUpdate, user=Depends(require_user)):
    if data.language not in ["en", "hi"]:
        raise HTTPException(400, "Supported: en, hi")
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("UPDATE users SET language=? WHERE id=?", (data.language, user["id"]))
        conn.commit()
    return {"message": "Language updated", "language": data.language}


# ============================================================
# AUDIT
# ============================================================
@app.get("/audit/logs")
def audit_logs(admin=Depends(require_admin), limit: int = Query(100, ge=1, le=1000)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM audit_log ORDER BY created_at DESC LIMIT ?", (limit,))
        return {"logs": [dict(r) for r in cursor.fetchall()]}


@app.get("/audit/login-attempts")
def login_attempts(admin=Depends(require_admin), limit: int = Query(100, ge=1, le=1000)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM login_attempts ORDER BY created_at DESC LIMIT ?", (limit,))
        return {"attempts": [dict(r) for r in cursor.fetchall()]}


# ============================================================
# PARENT-CHILD
# ============================================================
@app.post("/parents/link-child")
def link_child(request: Request, data: LinkChildRequest, admin=Depends(require_admin)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT id FROM users WHERE id=? AND role='parent'",
                       (data.parent_user_id,))
        if not cursor.fetchone():
            raise HTTPException(404, "Parent user not found")
        cursor.execute("SELECT id FROM students WHERE id=?", (data.student_id,))
        if not cursor.fetchone():
            raise HTTPException(404, "Student not found")

        try:
            cursor.execute("""
                INSERT INTO parent_children (parent_user_id, student_id, relationship)
                VALUES (?, ?, ?)
            """, (data.parent_user_id, data.student_id, data.relationship))
            conn.commit()
        except sqlite3.IntegrityError:
            raise HTTPException(400, "Already linked")

    return {"message": "Child linked successfully"}


@app.get("/parents/my-children")
def get_my_children(user=Depends(require_user)):
    if user["role"] != "parent":
        raise HTTPException(403, "Parents only")
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT s.id, s.name, s.grade, s.average, s.department,
                   s.semester, s.class_name, pc.relationship
            FROM students s
            JOIN parent_children pc ON s.id = pc.student_id
            WHERE pc.parent_user_id = ?
        """, (user["id"],))
        return {"children": [dict(r) for r in cursor.fetchall()]}


@app.get("/parents/all-links")
def get_all_links(admin=Depends(require_admin)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT pc.id, pc.relationship,
                   u.username AS parent_username, u.email AS parent_email,
                   s.name AS student_name, s.id AS student_id, u.id AS parent_id
            FROM parent_children pc
            JOIN users u ON pc.parent_user_id = u.id
            JOIN students s ON pc.student_id = s.id
            ORDER BY pc.created_at DESC
        """)
        return {"links": [dict(r) for r in cursor.fetchall()]}


@app.delete("/parents/unlink/{link_id}")
def unlink_child(request: Request, link_id: int, admin=Depends(require_admin)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM parent_children WHERE id=?", (link_id,))
        if cursor.rowcount == 0:
            raise HTTPException(404, "Link not found")
        conn.commit()
    return {"message": "Unlinked successfully"}
# ============================================================
# ACCESS HELPER
# ============================================================
def get_accessible_students(user):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        if user["role"] in ["admin", "teacher"]:
            cursor.execute("SELECT * FROM students ORDER BY id DESC")
        elif user["role"] == "parent":
            cursor.execute("""
                SELECT s.* FROM students s
                JOIN parent_children pc ON s.id = pc.student_id
                WHERE pc.parent_user_id = ? ORDER BY s.id DESC
            """, (user["id"],))
        elif user["role"] == "student":
            cursor.execute("SELECT * FROM students WHERE user_id = ? ORDER BY id DESC",
                           (user["id"],))
        else:
            return []

        students = []
        for r in cursor.fetchall():
            s = dict(r)
            s["marks"] = json.loads(s["marks"])
            s["subjects"] = json.loads(s["subjects"])
            students.append(s)
        return students


# ============================================================
# STUDENTS
# ============================================================
@app.get("/students")
def get_all_students(user=Depends(require_user),
                     limit: int = Query(100, ge=1, le=1000),
                     offset: int = Query(0, ge=0)):
    try:
        all_students = get_accessible_students(user)
        total = len(all_students)
        return {
            "count": total, "limit": limit, "offset": offset,
            "page": (offset // limit) + 1 if limit > 0 else 1,
            "total_pages": (total + limit - 1) // limit if limit > 0 else 1,
            "students": all_students[offset:offset + limit],
        }
    except Exception as e:
        raise HTTPException(500, f"Error: {str(e)}")


@app.get("/students/{student_id}")
def get_student_by_id(student_id: int, user=Depends(require_user)):
    for s in get_accessible_students(user):
        if s["id"] == student_id:
            return s
    raise HTTPException(404, "Not found or access denied")


@app.get("/students/{student_id}/profile")
def student_profile(student_id: int, user=Depends(require_user)):
    accessible = get_accessible_students(user)
    student = next((s for s in accessible if s["id"] == student_id), None)
    if not student:
        raise HTTPException(404, "Not found or access denied")

    with get_db_connection() as conn:
        cursor = conn.cursor()

        cursor.execute("""
            SELECT status, COUNT(*) as count FROM attendance
            WHERE student_id=? GROUP BY status
        """, (student_id,))
        att_stats = {r["status"]: r["count"] for r in cursor.fetchall()}

        cursor.execute("""
            SELECT * FROM performance_trends WHERE student_id=?
            ORDER BY created_at ASC
        """, (student_id,))
        trends = [dict(r) for r in cursor.fetchall()]

        cursor.execute("""
            SELECT * FROM notifications WHERE student_id=?
            ORDER BY created_at DESC LIMIT 20
        """, (student_id,))
        notifications = [dict(r) for r in cursor.fetchall()]

        cursor.execute("""
            SELECT u.username, u.email, u.full_name, pc.relationship
            FROM parent_children pc
            JOIN users u ON pc.parent_user_id = u.id
            WHERE pc.student_id=?
        """, (student_id,))
        parents = [dict(r) for r in cursor.fetchall()]

    return {
        "student": student,
        "attendance": att_stats,
        "trends": trends,
        "notifications": notifications,
        "parents": parents,
    }


@app.post("/students/save")
def save_student(student: StudentRecord, background_tasks: BackgroundTasks,
                 user=Depends(require_role("admin", "teacher"))):
    try:
        total = student.total_marks if student.total_marks is not None else sum(student.marks)
        with get_db_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                INSERT INTO students (name, marks, subjects, grade, average, total_marks,
                                      timestamp, semester, batch_year, department, class_name)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (student.name, json.dumps(student.marks), json.dumps(student.subjects),
                  student.grade, student.average, total,
                  student.timestamp or datetime.now().isoformat(),
                  student.semester, student.batch_year, student.department, student.class_name))
            sid = cursor.lastrowid
            cursor.execute("""
                INSERT INTO performance_trends (student_id, semester, average, grade,
                                                total_marks, timestamp)
                VALUES (?, ?, ?, ?, ?, ?)
            """, (sid, student.semester, student.average, student.grade, total,
                  datetime.now().isoformat()))
            if student.grade in ["A", "A+"]:
                cursor.execute(
                    "INSERT INTO notifications (student_id, message, type) VALUES (?, ?, ?)",
                    (sid, f"Congratulations! Grade {student.grade}", "Achievement"))
            elif student.grade == "F":
                cursor.execute(
                    "INSERT INTO notifications (student_id, message, type) VALUES (?, ?, ?)",
                    (sid, "Warning: Grade F. Seek support.", "Warning"))
            conn.commit()

        parent_emails = get_parent_emails_for_student(sid)
        if parent_emails:
            background_tasks.add_task(
                send_grade_notification_email,
                student.name, student.grade, student.average, parent_emails
            )

        return {"message": "Saved", "id": sid}
    except Exception as e:
        raise HTTPException(500, f"Error: {str(e)}")


@app.delete("/students/delete/{student_id}")
def delete_student(student_id: int, user=Depends(require_role("admin", "teacher"))):
    try:
        with get_db_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT id FROM students WHERE id=?", (student_id,))
            if not cursor.fetchone():
                raise HTTPException(404, "Not found")
            cursor.execute("DELETE FROM attendance WHERE student_id=?", (student_id,))
            cursor.execute("DELETE FROM notifications WHERE student_id=?", (student_id,))
            cursor.execute("DELETE FROM performance_trends WHERE student_id=?", (student_id,))
            cursor.execute("DELETE FROM parent_children WHERE student_id=?", (student_id,))
            cursor.execute("DELETE FROM students WHERE id=?", (student_id,))
            conn.commit()
        return {"message": "Deleted"}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"Error: {str(e)}")


# ============================================================
# PHASE 3: Student Photo Upload
# ============================================================
@app.post("/students/{student_id}/photo")
async def upload_student_photo(
    student_id: int,
    file: UploadFile = File(...),
    user=Depends(require_role("admin", "teacher")),
):
    accessible = get_accessible_students(user)
    if not any(s["id"] == student_id for s in accessible):
        raise HTTPException(404, "Student not found")

    if file.content_type not in ["image/jpeg", "image/png", "image/webp", "image/jpg"]:
        raise HTTPException(400, "Only JPG, PNG, WEBP images allowed")

    content = await file.read()
    if len(content) > MAX_UPLOAD_SIZE_MB * 1024 * 1024:
        raise HTTPException(400, f"File too large (max {MAX_UPLOAD_SIZE_MB} MB)")

    ext = os.path.splitext(file.filename or "photo.jpg")[1].lower() or ".jpg"
    filename = f"student_{student_id}_{int(datetime.now().timestamp())}{ext}"
    filepath = os.path.join(UPLOAD_DIR, "photos", filename)

    with open(filepath, "wb") as f:
        f.write(content)

    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "UPDATE students SET photo_url=? WHERE id=?",
            (f"/uploads/photos/{filename}", student_id),
        )
        conn.commit()

    log_live_event(user["id"], "photo_uploaded", {
        "student_id": student_id,
        "photo_url": f"/uploads/photos/{filename}",
    })

    return {
        "message": "Photo uploaded",
        "photo_url": f"/uploads/photos/{filename}",
    }


@app.delete("/students/{student_id}/photo")
def delete_student_photo(
    student_id: int,
    user=Depends(require_role("admin", "teacher")),
):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT photo_url FROM students WHERE id=?", (student_id,))
        row = cursor.fetchone()
        if not row:
            raise HTTPException(404, "Student not found")

        if row["photo_url"]:
            old_file = row["photo_url"].replace("/uploads/", "")
            old_path = os.path.join(UPLOAD_DIR, old_file)
            if os.path.exists(old_path):
                try:
                    os.remove(old_path)
                except Exception:
                    pass

        cursor.execute("UPDATE students SET photo_url=NULL WHERE id=?", (student_id,))
        conn.commit()

    return {"message": "Photo deleted"}


@app.put("/students/{student_id}/profile")
def update_student_profile(
    student_id: int,
    data: StudentProfileUpdate,
    user=Depends(require_role("admin", "teacher")),
):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT id FROM students WHERE id=?", (student_id,))
        if not cursor.fetchone():
            raise HTTPException(404, "Student not found")

        update_data = data.dict(exclude_unset=True)
        if not update_data:
            raise HTTPException(400, "No fields to update")

        set_clause = ", ".join(f"{k}=?" for k in update_data.keys())
        values = list(update_data.values()) + [student_id]

        cursor.execute(f"UPDATE students SET {set_clause} WHERE id=?", values)
        conn.commit()

    return {"message": "Profile updated"}


# ============================================================
# PHASE 3: Attendance Heatmap
# ============================================================
@app.get("/attendance/{student_id}/heatmap")
def attendance_heatmap(
    student_id: int,
    year: Optional[int] = None,
    user=Depends(require_user),
):
    if not any(s["id"] == student_id for s in get_accessible_students(user)):
        raise HTTPException(403, "Access denied")

    if not year:
        year = datetime.now().year

    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT date, status, COUNT(*) as count
            FROM attendance
            WHERE student_id=? AND date LIKE ?
            GROUP BY date, status
        """, (student_id, f"{year}%"))

        rows = cursor.fetchall()

        by_date = {}
        for r in rows:
            d = r["date"]
            if d not in by_date:
                by_date[d] = {"present": 0, "absent": 0, "late": 0, "excused": 0}
            status_key = r["status"].lower()
            if status_key in by_date[d]:
                by_date[d][status_key] += r["count"]

        heatmap = []
        for date_str, counts in by_date.items():
            total = sum(counts.values())
            if total == 0:
                continue
            score = (counts["present"] + counts["late"] * 0.5 + counts["excused"] * 0.5) / total
            heatmap.append({
                "date": date_str,
                "count": total,
                "score": round(score, 2),
                "status": "present" if score >= 0.75 else "partial" if score >= 0.5 else "absent",
            })

        return {
            "student_id": student_id,
            "year": year,
            "heatmap": sorted(heatmap, key=lambda x: x["date"]),
        }


# ============================================================
# PHASE 3: EXAMS
# ============================================================
@app.post("/exams/create")
def create_exam(
    data: ExamCreate,
    user=Depends(require_role("admin", "teacher")),
):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO exams (name, class_name, subject, exam_date, start_time, end_time,
                              total_marks, room, notes)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (data.name, data.class_name, data.subject, data.exam_date,
              data.start_time, data.end_time, data.total_marks, data.room, data.notes))
        conn.commit()
        exam_id = cursor.lastrowid

    log_live_event(user["id"], "exam_created", {"exam_id": exam_id, "name": data.name})
    return {"message": "Exam created", "id": exam_id}


@app.get("/exams")
def list_exams(
    class_name: Optional[str] = None,
    user=Depends(require_user),
):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        if class_name:
            cursor.execute(
                "SELECT * FROM exams WHERE class_name=? ORDER BY exam_date ASC",
                (class_name,),
            )
        else:
            cursor.execute("SELECT * FROM exams ORDER BY exam_date ASC")
        return {"exams": [dict(r) for r in cursor.fetchall()]}


@app.delete("/exams/{exam_id}")
def delete_exam(exam_id: int, user=Depends(require_role("admin", "teacher"))):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM exams WHERE id=?", (exam_id,))
        if cursor.rowcount == 0:
            raise HTTPException(404, "Exam not found")
        conn.commit()
    return {"message": "Exam deleted"}


# ============================================================
# PHASE 3: TIMETABLE
# ============================================================
@app.post("/timetable/create")
def create_timetable_entry(
    data: TimetableCreate,
    user=Depends(require_role("admin", "teacher")),
):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO timetable (class_name, day_of_week, period, subject,
                                   teacher_name, room, start_time, end_time)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """, (data.class_name, data.day_of_week, data.period, data.subject,
              data.teacher_name, data.room, data.start_time, data.end_time))
        conn.commit()
    return {"message": "Timetable entry created", "id": cursor.lastrowid}


@app.get("/timetable")
def list_timetable(
    class_name: Optional[str] = None,
    user=Depends(require_user),
):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        if class_name:
            cursor.execute(
                """SELECT * FROM timetable WHERE class_name=?
                   ORDER BY day_of_week, period""",
                (class_name,),
            )
        else:
            cursor.execute("SELECT * FROM timetable ORDER BY class_name, day_of_week, period")
        return {"timetable": [dict(r) for r in cursor.fetchall()]}


@app.delete("/timetable/{entry_id}")
def delete_timetable_entry(entry_id: int, user=Depends(require_role("admin", "teacher"))):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM timetable WHERE id=?", (entry_id,))
        if cursor.rowcount == 0:
            raise HTTPException(404, "Entry not found")
        conn.commit()
    return {"message": "Entry deleted"}


# ============================================================
# PHASE 3: ASSIGNMENTS
# ============================================================
@app.post("/assignments/create")
def create_assignment(
    data: AssignmentCreate,
    user=Depends(require_role("admin", "teacher")),
):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO assignments (title, description, class_name, subject,
                                    due_date, total_marks, created_by)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (data.title, data.description, data.class_name, data.subject,
              data.due_date, data.total_marks, user["id"]))
        conn.commit()
        aid = cursor.lastrowid

    log_live_event(user["id"], "assignment_created", {"assignment_id": aid})
    return {"message": "Assignment created", "id": aid}


@app.get("/assignments")
def list_assignments(
    class_name: Optional[str] = None,
    user=Depends(require_user),
):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        if class_name:
            cursor.execute(
                "SELECT * FROM assignments WHERE class_name=? ORDER BY due_date ASC",
                (class_name,),
            )
        else:
            cursor.execute("SELECT * FROM assignments ORDER BY due_date ASC")

        assignments = [dict(r) for r in cursor.fetchall()]

        for a in assignments:
            cursor.execute("""
                SELECT COUNT(*) as total,
                       SUM(CASE WHEN status='submitted' THEN 1 ELSE 0 END) as submitted,
                       SUM(CASE WHEN status='graded' THEN 1 ELSE 0 END) as graded
                FROM assignment_submissions WHERE assignment_id=?
            """, (a["id"],))
            row = cursor.fetchone()
            a["stats"] = {
                "total": row["total"] or 0,
                "submitted": row["submitted"] or 0,
                "graded": row["graded"] or 0,
            }

        return {"assignments": assignments}


@app.get("/assignments/{assignment_id}/submissions")
def list_submissions(
    assignment_id: int,
    user=Depends(require_role("admin", "teacher")),
):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT s.id as student_id, s.name as student_name,
                   sub.id as submission_id, sub.status, sub.marks_obtained,
                   sub.feedback, sub.submitted_at
            FROM students s
            LEFT JOIN assignment_submissions sub
                ON sub.student_id = s.id AND sub.assignment_id = ?
            WHERE s.class_name = (SELECT class_name FROM assignments WHERE id=?)
            ORDER BY s.name
        """, (assignment_id, assignment_id))
        return {"submissions": [dict(r) for r in cursor.fetchall()]}


@app.put("/assignments/submissions/{submission_id}")
def update_submission(
    submission_id: int,
    data: SubmissionUpdate,
    user=Depends(require_role("admin", "teacher")),
):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        update_data = data.dict(exclude_unset=True)
        if not update_data:
            raise HTTPException(400, "No fields to update")

        set_clause = ", ".join(f"{k}=?" for k in update_data.keys())
        values = list(update_data.values()) + [submission_id]

        cursor.execute(f"UPDATE assignment_submissions SET {set_clause} WHERE id=?", values)
        conn.commit()

    return {"message": "Submission updated"}


@app.post("/assignments/{assignment_id}/submit/{student_id}")
def submit_assignment(
    assignment_id: int,
    student_id: int,
    user=Depends(require_role("admin", "teacher")),
):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO assignment_submissions (assignment_id, student_id, status, submitted_at)
            VALUES (?, ?, 'submitted', ?)
            ON CONFLICT(assignment_id, student_id)
            DO UPDATE SET status='submitted', submitted_at=?
        """, (assignment_id, student_id, datetime.now().isoformat(), datetime.now().isoformat()))
        conn.commit()

    return {"message": "Marked as submitted"}


@app.delete("/assignments/{assignment_id}")
def delete_assignment(assignment_id: int, user=Depends(require_role("admin", "teacher"))):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM assignment_submissions WHERE assignment_id=?", (assignment_id,))
        cursor.execute("DELETE FROM assignments WHERE id=?", (assignment_id,))
        if cursor.rowcount == 0:
            raise HTTPException(404, "Assignment not found")
        conn.commit()
    return {"message": "Assignment deleted"}


# ============================================================
# PHASE 3: FEES
# ============================================================
@app.post("/fees/structure/create")
def create_fee_structure(
    data: FeeStructureCreate,
    user=Depends(require_admin),
):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO fee_structures (class_name, fee_type, amount, frequency, academic_year)
            VALUES (?, ?, ?, ?, ?)
        """, (data.class_name, data.fee_type, data.amount, data.frequency, data.academic_year))
        conn.commit()
    return {"message": "Fee structure created", "id": cursor.lastrowid}


@app.get("/fees/structure")
def list_fee_structures(user=Depends(require_user)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM fee_structures ORDER BY class_name, fee_type")
        return {"structures": [dict(r) for r in cursor.fetchall()]}


@app.delete("/fees/structure/{structure_id}")
def delete_fee_structure(structure_id: int, user=Depends(require_admin)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM fee_structures WHERE id=?", (structure_id,))
        conn.commit()
    return {"message": "Fee structure deleted"}


@app.post("/fees/payment/create")
def record_fee_payment(
    data: FeePaymentCreate,
    user=Depends(require_role("admin", "teacher")),
):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO fee_payments (student_id, fee_type, amount, payment_date,
                                     payment_method, transaction_id, status, due_date, notes)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (data.student_id, data.fee_type, data.amount, data.payment_date,
              data.payment_method, data.transaction_id, data.status,
              data.due_date, data.notes))
        conn.commit()
        payment_id = cursor.lastrowid

    log_live_event(user["id"], "fee_paid", {
        "student_id": data.student_id,
        "amount": data.amount,
    })
    return {"message": "Payment recorded", "id": payment_id}


@app.get("/fees/payments")
def list_fee_payments(
    student_id: Optional[int] = None,
    user=Depends(require_user),
):
    with get_db_connection() as conn:
        cursor = conn.cursor()

        if user["role"] == "student":
            cursor.execute("SELECT id FROM students WHERE user_id=?", (user["id"],))
            row = cursor.fetchone()
            if not row:
                return {"payments": [], "summary": {}}
            student_id = row["id"]

        if student_id:
            cursor.execute("""
                SELECT fp.*, s.name as student_name
                FROM fee_payments fp
                JOIN students s ON fp.student_id = s.id
                WHERE fp.student_id=?
                ORDER BY fp.payment_date DESC
            """, (student_id,))
        else:
            cursor.execute("""
                SELECT fp.*, s.name as student_name
                FROM fee_payments fp
                JOIN students s ON fp.student_id = s.id
                ORDER BY fp.payment_date DESC
                LIMIT 500
            """)

        payments = [dict(r) for r in cursor.fetchall()]

        total_paid = sum(p["amount"] for p in payments if p["status"] == "paid")
        total_pending = sum(p["amount"] for p in payments if p["status"] == "pending")

        return {
            "payments": payments,
            "summary": {
                "total_paid": total_paid,
                "total_pending": total_pending,
                "count": len(payments),
            },
        }


@app.get("/fees/student/{student_id}/summary")
def student_fee_summary(student_id: int, user=Depends(require_user)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT fee_type,
                   SUM(CASE WHEN status='paid' THEN amount ELSE 0 END) as paid,
                   SUM(CASE WHEN status='pending' THEN amount ELSE 0 END) as pending
            FROM fee_payments WHERE student_id=?
            GROUP BY fee_type
        """, (student_id,))
        summary = [dict(r) for r in cursor.fetchall()]
    return {"summary": summary}


# ============================================================
# PHASE 3: SAVED FILTERS
# ============================================================
@app.post("/filters/save")
def save_filter(data: SavedFilterCreate, user=Depends(require_user)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO saved_filters (user_id, name, entity, filter_json, is_shared)
            VALUES (?, ?, ?, ?, ?)
        """, (user["id"], data.name, data.entity, data.filter_json,
              1 if data.is_shared else 0))
        conn.commit()
    return {"message": "Filter saved", "id": cursor.lastrowid}


@app.get("/filters")
def list_saved_filters(user=Depends(require_user)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT * FROM saved_filters
            WHERE user_id=? OR is_shared=1
            ORDER BY created_at DESC
        """, (user["id"],))
        return {"filters": [dict(r) for r in cursor.fetchall()]}


@app.delete("/filters/{filter_id}")
def delete_saved_filter(filter_id: int, user=Depends(require_user)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "DELETE FROM saved_filters WHERE id=? AND user_id=?",
            (filter_id, user["id"]),
        )
        if cursor.rowcount == 0:
            raise HTTPException(404, "Filter not found or not yours")
        conn.commit()
    return {"message": "Filter deleted"}


# ============================================================
# PHASE 3: BACKUP & RESTORE
# ============================================================
@app.post("/backup/create")
def create_backup(user=Depends(require_admin)):
    try:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_filename = f"backup_{timestamp}.zip"
        backup_path = os.path.join(BACKUP_DIR, backup_filename)

        with zipfile.ZipFile(backup_path, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.write(DB_PATH, "student_marks.db")

            if os.path.exists(UPLOAD_DIR):
                for root, dirs, files in os.walk(UPLOAD_DIR):
                    for file in files:
                        filepath = os.path.join(root, file)
                        arcname = os.path.relpath(filepath, os.path.dirname(UPLOAD_DIR))
                        zf.write(filepath, arcname)

        size = os.path.getsize(backup_path)

        with get_db_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                INSERT INTO backup_logs (filename, size_bytes, backup_type, created_by)
                VALUES (?, ?, ?, ?)
            """, (backup_filename, size, "manual", user["id"]))
            conn.commit()

        return {
            "message": "Backup created",
            "filename": backup_filename,
            "size_mb": round(size / 1024 / 1024, 2),
        }
    except Exception as e:
        raise HTTPException(500, f"Backup failed: {str(e)}")


@app.get("/backup/list")
def list_backups(user=Depends(require_admin)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM backup_logs ORDER BY created_at DESC LIMIT 50")
        return {"backups": [dict(r) for r in cursor.fetchall()]}


@app.get("/backup/download/{filename}")
def download_backup(filename: str, user=Depends(require_admin)):
    filepath = os.path.join(BACKUP_DIR, filename)
    if not os.path.exists(filepath):
        raise HTTPException(404, "Backup not found")
    return FileResponse(filepath, filename=filename, media_type="application/zip")


@app.post("/backup/restore")
async def restore_backup(
    file: UploadFile = File(...),
    user=Depends(require_admin),
):
    try:
        content = await file.read()

        temp_path = os.path.join(BACKUP_DIR, f"restore_temp_{int(datetime.now().timestamp())}.zip")
        with open(temp_path, "wb") as f:
            f.write(content)

        if not zipfile.is_zipfile(temp_path):
            os.remove(temp_path)
            raise HTTPException(400, "Invalid ZIP file")

        emergency_backup = os.path.join(BACKUP_DIR, f"emergency_{int(datetime.now().timestamp())}.db")
        shutil.copy(DB_PATH, emergency_backup)

        extract_dir = os.path.join(BACKUP_DIR, "restore_temp")
        os.makedirs(extract_dir, exist_ok=True)

        with zipfile.ZipFile(temp_path, "r") as zf:
            zf.extractall(extract_dir)

        restored_db = os.path.join(extract_dir, "student_marks.db")
        if os.path.exists(restored_db):
            shutil.copy(restored_db, DB_PATH)

        shutil.rmtree(extract_dir, ignore_errors=True)
        os.remove(temp_path)

        return {
            "message": "Backup restored successfully",
            "emergency_backup": os.path.basename(emergency_backup),
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"Restore failed: {str(e)}")


@app.get("/backup/auto-list")
def list_auto_backups(user=Depends(require_admin)):
    auto_dir = os.path.join(BACKUP_DIR, "auto")
    if not os.path.exists(auto_dir):
        return {"backups": []}

    backups = []
    for f in sorted(os.listdir(auto_dir), reverse=True):
        if f.endswith(".zip"):
            path = os.path.join(auto_dir, f)
            backups.append({
                "filename": f,
                "size_mb": round(os.path.getsize(path) / 1024 / 1024, 2),
                "created_at": datetime.fromtimestamp(os.path.getmtime(path)).isoformat(),
            })
    return {"backups": backups}
# ============================================================
# PHASE 3: SCHEDULED REPORTS
# ============================================================
@app.post("/reports/schedule")
def create_scheduled_report(
    data: ScheduledReportCreate,
    user=Depends(require_admin),
):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            INSERT INTO scheduled_reports (report_type, recipients, schedule, enabled)
            VALUES (?, ?, ?, ?)
        """, (data.report_type, data.recipients, data.schedule,
              1 if data.enabled else 0))
        conn.commit()
    return {"message": "Scheduled report created", "id": cursor.lastrowid}


@app.get("/reports/scheduled")
def list_scheduled_reports(user=Depends(require_admin)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM scheduled_reports ORDER BY created_at DESC")
        return {"reports": [dict(r) for r in cursor.fetchall()]}


@app.delete("/reports/scheduled/{report_id}")
def delete_scheduled_report(report_id: int, user=Depends(require_admin)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM scheduled_reports WHERE id=?", (report_id,))
        conn.commit()
    return {"message": "Scheduled report deleted"}


def send_weekly_report():
    """Background task — sends weekly reports."""
    try:
        with get_db_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM students")
            students = [dict(r) for r in cursor.fetchall()]

            cursor.execute("SELECT * FROM scheduled_reports WHERE enabled=1")
            reports = [dict(r) for r in cursor.fetchall()]

        if not students or not reports:
            return

        avg = sum(s["average"] for s in students) / len(students)
        html = f"""
        <html><body style="font-family: Arial, sans-serif; padding: 20px;">
            <h1 style="color: #667eea;">Weekly Report</h1>
            <p><b>Total Students:</b> {len(students)}</p>
            <p><b>Average Score:</b> {avg:.2f}%</p>
            <p><b>Date:</b> {datetime.now().strftime('%Y-%m-%d')}</p>
            <p>Log in to Student Marks Analyzer for full details.</p>
        </body></html>
        """

        for report in reports:
            recipients = [r.strip() for r in report["recipients"].split(",") if r.strip()]
            for email in recipients:
                send_email_notification(email, "Weekly Performance Report", html)
                print(f"[REPORT] Sent to {email}")

        with get_db_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(
                "UPDATE scheduled_reports SET last_sent=? WHERE enabled=1",
                (datetime.now().isoformat(),),
            )
            conn.commit()
    except Exception as e:
        print(f"[REPORT ERROR] {e}")


# ============================================================
# PHASE 3: BULK OPERATIONS (existing + extended)
# ============================================================
@app.post("/students/bulk-delete")
def bulk_delete_students(request: Request, data: BulkDeleteRequest,
                         user=Depends(require_role("admin", "teacher"))):
    if not data.student_ids:
        raise HTTPException(400, "No student IDs provided")

    deleted = 0
    with get_db_connection() as conn:
        cursor = conn.cursor()
        for sid in data.student_ids:
            cursor.execute("DELETE FROM attendance WHERE student_id=?", (sid,))
            cursor.execute("DELETE FROM notifications WHERE student_id=?", (sid,))
            cursor.execute("DELETE FROM performance_trends WHERE student_id=?", (sid,))
            cursor.execute("DELETE FROM parent_children WHERE student_id=?", (sid,))
            cursor.execute("DELETE FROM assignment_submissions WHERE student_id=?", (sid,))
            cursor.execute("DELETE FROM fee_payments WHERE student_id=?", (sid,))
            cursor.execute("DELETE FROM students WHERE id=?", (sid,))
            deleted += cursor.rowcount
        conn.commit()

    return {"message": f"Deleted {deleted} students", "deleted": deleted}


@app.post("/notifications/bulk-send")
def bulk_send_notifications(request: Request, data: BulkNotificationRequest,
                            user=Depends(require_role("admin", "teacher"))):
    if not data.student_ids:
        raise HTTPException(400, "No student IDs provided")

    sent = 0
    with get_db_connection() as conn:
        cursor = conn.cursor()
        for sid in data.student_ids:
            cursor.execute(
                "INSERT INTO notifications (student_id, message, type) VALUES (?, ?, ?)",
                (sid, data.message, data.type)
            )
            sent += 1
        conn.commit()

    return {"message": f"Sent {sent} notifications", "sent": sent}


@app.post("/attendance/bulk")
def bulk_attendance(request: Request, data: BulkAttendanceRequest,
                    user=Depends(require_role("admin", "teacher"))):
    if not data.student_ids:
        raise HTTPException(400, "No student IDs")

    recorded = 0
    with get_db_connection() as conn:
        cursor = conn.cursor()
        for sid in data.student_ids:
            cursor.execute("""
                INSERT INTO attendance (student_id, date, status, subject)
                VALUES (?, ?, ?, ?)
            """, (sid, data.date, data.status, data.subject))
            if data.status in ["Absent", "Late"]:
                cursor.execute(
                    "INSERT INTO notifications (student_id, message, type) VALUES (?, ?, ?)",
                    (sid, f"Marked {data.status} on {data.date}", "Warning"))
            recorded += 1
        conn.commit()

    return {"message": f"Recorded attendance for {recorded} students",
            "recorded": recorded}


# ============================================================
# CLASSES (existing)
# ============================================================
@app.post("/classes/create")
def create_class(request: Request, data: ClassCreate, admin=Depends(require_admin)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        try:
            cursor.execute("""
                INSERT INTO classes (name, department, teacher_id)
                VALUES (?, ?, ?)
            """, (data.name, data.department, data.teacher_id))
            conn.commit()
        except sqlite3.IntegrityError:
            raise HTTPException(400, "Class already exists")
    return {"message": "Class created", "id": cursor.lastrowid}


@app.get("/classes")
def list_classes(user=Depends(require_user)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM classes ORDER BY name")
        return {"classes": [dict(r) for r in cursor.fetchall()]}


@app.delete("/classes/{class_id}")
def delete_class(class_id: int, admin=Depends(require_admin)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM classes WHERE id=?", (class_id,))
        if cursor.rowcount == 0:
            raise HTTPException(404, "Class not found")
        conn.commit()
    return {"message": "Class deleted"}


@app.post("/classes/assign")
def assign_students(data: ClassAssignRequest, admin=Depends(require_admin)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        updated = 0
        for sid in data.student_ids:
            cursor.execute("UPDATE students SET class_name=? WHERE id=?",
                           (data.class_name, sid))
            updated += cursor.rowcount
        conn.commit()
    return {"message": f"Assigned {updated} students to {data.class_name}",
            "updated": updated}


@app.get("/classes/{class_name}/analytics")
def class_analytics(class_name: str, user=Depends(require_user)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM students WHERE class_name=?", (class_name,))
        students = []
        for r in cursor.fetchall():
            s = dict(r)
            s["marks"] = json.loads(s["marks"])
            s["subjects"] = json.loads(s["subjects"])
            students.append(s)

        if not students:
            return {"message": f"No students in class '{class_name}'",
                    "class_name": class_name, "count": 0}

        avgs = [s["average"] for s in students]
        grades = [s["grade"] for s in students]

        dist = {}
        for g in grades:
            dist[g] = dist.get(g, 0) + 1

        return {
            "class_name": class_name,
            "count": len(students),
            "average": round(float(np.mean(avgs)), 2),
            "highest": round(float(max(avgs)), 2),
            "lowest": round(float(min(avgs)), 2),
            "std_dev": round(float(np.std(avgs)), 2),
            "median": round(float(np.median(avgs)), 2),
            "grade_distribution": dist,
            "pass_rate": round(sum(1 for g in grades if g != "F") / len(students) * 100, 2),
            "top_3": sorted([
                {"id": s["id"], "name": s["name"], "average": s["average"], "grade": s["grade"]}
                for s in students
            ], key=lambda x: x["average"], reverse=True)[:3],
            "bottom_3": sorted([
                {"id": s["id"], "name": s["name"], "average": s["average"], "grade": s["grade"]}
                for s in students
            ], key=lambda x: x["average"])[:3],
            "at_risk": [
                {"id": s["id"], "name": s["name"], "average": s["average"]}
                for s in students if s["average"] < 50
            ],
        }


@app.get("/classes/list-names")
def class_names(user=Depends(require_user)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT DISTINCT class_name FROM students WHERE class_name IS NOT NULL")
        from_students = [r["class_name"] for r in cursor.fetchall()]
        cursor.execute("SELECT name FROM classes")
        from_table = [r["name"] for r in cursor.fetchall()]
        all_names = list(set(from_students + from_table))
        return {"names": sorted(all_names)}


# ============================================================
# EXCEL EXPORT
# ============================================================
@app.get("/export/excel")
def export_excel(user=Depends(require_user)):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        students = get_accessible_students(user)
        if not students:
            raise HTTPException(404, "No data")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Students"

        headers = ["ID", "Name", "Marks", "Subjects", "Grade", "Average",
                   "Total", "Semester", "Department", "Class", "Batch Year", "Created"]
        ws.append(headers)

        header_fill = PatternFill(start_color="667EEA", end_color="667EEA", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for s in students:
            ws.append([
                s["id"], s["name"],
                ", ".join(str(m) for m in s["marks"]),
                ", ".join(s["subjects"]),
                s["grade"], s["average"], s["total_marks"],
                s.get("semester") or "",
                s.get("department") or "",
                s.get("class_name") or "",
                s.get("batch_year") or "",
                s.get("created_at") or "",
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition":
                     f"attachment; filename=students_{datetime.now().strftime('%Y%m%d')}.xlsx"}
        )
    except ImportError:
        raise HTTPException(500, "openpyxl not installed")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"Excel export error: {str(e)}")


# ============================================================
# CSV IMPORT/EXPORT
# ============================================================
@app.get("/students/import-template")
def get_import_template(user=Depends(require_role("admin", "teacher"))):
    template = """name,marks,subjects,semester,department,batch_year,class_name
John Doe,"85,92,78,88,91","Math,Science,English,History,Art",Fall 2024,Computer Science,2024,CS-A
"""
    return Response(content=template, media_type="text/csv",
                    headers={"Content-Disposition": "attachment; filename=template.csv"})


@app.post("/students/import-csv")
async def import_csv(user=Depends(require_role("admin", "teacher")),
                     file: UploadFile = File(...)):
    try:
        content = await file.read()
        text = None
        for enc in ["utf-8", "utf-8-sig", "latin-1"]:
            try:
                text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise HTTPException(400, "Cannot decode file")

        reader = csv_module.DictReader(io.StringIO(text))
        fmap = {f.lower().strip(): f for f in (reader.fieldnames or [])}
        missing = [r for r in ["name", "marks", "subjects"] if r not in fmap]
        if missing:
            raise HTTPException(400, f"Missing columns: {missing}")

        imported = 0
        failed = []
        row_num = 1
        with get_db_connection() as conn:
            cursor = conn.cursor()
            for row in reader:
                row_num += 1
                try:
                    name = str(row.get(fmap["name"], "")).strip()
                    if not name:
                        failed.append(f"Row {row_num}: empty name")
                        continue
                    marks_raw = str(row.get(fmap["marks"], "")).strip()
                    marks = []
                    for sep in [",", ";", "|"]:
                        if sep in marks_raw:
                            marks = [float(m.strip()) for m in marks_raw.split(sep) if m.strip()]
                            break
                    if not marks:
                        marks = [float(marks_raw)]

                    subj_raw = str(row.get(fmap["subjects"], "")).strip()
                    subjects = []
                    for sep in [",", ";", "|"]:
                        if sep in subj_raw:
                            subjects = [s.strip() for s in subj_raw.split(sep) if s.strip()]
                            break
                    if not subjects:
                        subjects = [subj_raw]

                    if len(marks) != len(subjects):
                        failed.append(f"Row {row_num}: marks/subjects mismatch")
                        continue
                    if any(m < 0 or m > 100 for m in marks):
                        failed.append(f"Row {row_num}: marks 0-100 only")
                        continue

                    for sname in subjects:
                        ok, err = validate_subject_name(sname)
                        if not ok:
                            failed.append(f"Row {row_num}: {err}")
                            break
                    else:
                        stats = calculate_statistics(marks)
                        grade, _ = determine_grade(stats["average"])
                        semester = str(row.get(fmap.get("semester", ""), "") or "").strip() or None
                        dept = str(row.get(fmap.get("department", ""), "") or "").strip() or None
                        by = str(row.get(fmap.get("batch_year", ""), "") or "").strip() or None
                        cn = str(row.get(fmap.get("class_name", ""), "") or "").strip() or None

                        cursor.execute("""
                            INSERT INTO students (name, marks, subjects, grade, average,
                                                  total_marks, timestamp, semester,
                                                  batch_year, department, class_name)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """, (name, json.dumps(marks), json.dumps(subjects), grade,
                              stats["average"], stats["total"], datetime.now().isoformat(),
                              semester, by, dept, cn))
                        imported += 1
                except Exception as e:
                    failed.append(f"Row {row_num}: {str(e)}")
            conn.commit()
        return {"message": f"Imported {imported}", "imported": imported,
                "failed_count": len(failed), "failed": failed[:20]}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"Error: {str(e)}")


@app.get("/export/csv")
def export_csv(user=Depends(require_user)):
    try:
        students = get_accessible_students(user)
        if not students:
            raise HTTPException(404, "No data")
        out = io.StringIO()
        out.write("id,name,marks,subjects,grade,average,total_marks,semester,department,batch_year,class_name\n")
        for s in students:
            marks_str = '"' + ",".join(str(m) for m in s["marks"]) + '"'
            subj_str = '"' + ",".join(s["subjects"]) + '"'
            out.write(f'{s["id"]},"{s["name"]}",{marks_str},{subj_str},{s["grade"]},'
                      f'{s["average"]},{s["total_marks"]},{s.get("semester") or ""},'
                      f'{s.get("department") or ""},{s.get("batch_year") or ""},{s.get("class_name") or ""}\n')
        return Response(content=out.getvalue(), media_type="text/csv",
                        headers={"Content-Disposition":
                                 f"attachment; filename=students_{datetime.now().strftime('%Y%m%d')}.csv"})
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"Error: {str(e)}")


@app.get("/export/json")
def export_json(user=Depends(require_user)):
    try:
        students = get_accessible_students(user)
        if not students:
            raise HTTPException(404, "No data")
        return {"export_date": datetime.now().isoformat(),
                "total_students": len(students), "students": students}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"Error: {str(e)}")


# ============================================================
# ANALYZE
# ============================================================
@app.post("/analyze")
def analyze_marks(request: MarksRequest, user=Depends(require_user)):
    marks = request.marks
    threshold = request.passing_threshold
    stats = calculate_statistics(marks)
    grade, gp = determine_grade(stats["average"], request.grade_scheme)
    passed = sum(1 for m in marks if m >= threshold)
    failed = len(marks) - passed

    subject_wise = []
    for i, mark in enumerate(marks):
        subj = (request.subject_names[i] if request.subject_names and i < len(request.subject_names)
                else f"Subject {i+1}")
        subject_wise.append({
            "subject": subj, "marks": mark,
            "status": "Pass" if mark >= threshold else "Fail",
            "performance": "Excellent" if mark >= 80 else "Good" if mark >= 60
                            else "Average" if mark >= 40 else "Needs Improvement",
        })

    return {
        "student_name": request.student_name or "Unnamed",
        "semester": request.semester, "batch_year": request.batch_year,
        "department": request.department, "class_name": request.class_name,
        "total_marks": stats["total"], "average": stats["average"],
        "highest": stats["highest"], "lowest": stats["lowest"],
        "median": stats["median"], "std_deviation": stats["std_deviation"],
        "variance": stats["variance"], "quartiles": stats["quartiles"],
        "passed": passed, "failed": failed,
        "pass_percentage": (passed / len(marks)) * 100,
        "grade": grade, "grade_points": gp,
        "subject_wise": subject_wise,
        "performance_summary": {
            "performance_level": "Excellent" if stats["average"] >= 80 else "Good"
                                 if stats["average"] >= 60 else "Average"
                                 if stats["average"] >= 40 else "Poor",
            "consistency": "High" if stats["std_deviation"] < 15
                           else "Medium" if stats["std_deviation"] < 30 else "Low",
        },
        "ai_insights": generate_ai_insights(marks, stats["average"], request.subject_names or []),
        "anomalies": detect_anomalies(marks),
        "recommendations": generate_recommendations(marks, stats["average"]),
        "timestamp": datetime.now().isoformat(),
    }


# ============================================================
# ATTENDANCE
# ============================================================
@app.post("/attendance")
def record_attendance(att: AttendanceRecord,
                     user=Depends(require_role("admin", "teacher"))):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT id, name FROM students WHERE id=?", (att.student_id,))
        s = cursor.fetchone()
        if not s:
            raise HTTPException(404, "Student not found")
        cursor.execute("""
            INSERT INTO attendance (student_id, date, status, subject, notes)
            VALUES (?, ?, ?, ?, ?)
        """, (att.student_id, att.date, att.status, att.subject, att.notes))
        if att.status in ["Absent", "Late"]:
            cursor.execute(
                "INSERT INTO notifications (student_id, message, type) VALUES (?, ?, ?)",
                (att.student_id, f"Marked {att.status} on {att.date}", "Warning"))
        conn.commit()
        new_id = cursor.lastrowid
    return {"message": "Recorded", "id": new_id}


@app.get("/attendance/stats/overall")
def attendance_overall(user=Depends(require_user)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) as t FROM attendance")
        total = cursor.fetchone()["t"]
        if total == 0:
            return {"total_records": 0, "message": "No records"}
        cursor.execute("SELECT status, COUNT(*) as c FROM attendance GROUP BY status")
        counts = {r["status"]: r["c"] for r in cursor.fetchall()}
        return {
            "total_records": total,
            "present": counts.get("Present", 0),
            "absent": counts.get("Absent", 0),
            "late": counts.get("Late", 0),
            "excused": counts.get("Excused", 0),
            "overall_rate": round(counts.get("Present", 0) / total * 100, 2),
        }


@app.get("/attendance/{student_id}/stats")
def attendance_stats(student_id: int, user=Depends(require_user)):
    if not any(s["id"] == student_id for s in get_accessible_students(user)):
        raise HTTPException(403, "Access denied")
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM students WHERE id=?", (student_id,))
        s = cursor.fetchone()
        if not s:
            raise HTTPException(404, "Student not found")
        cursor.execute(
            "SELECT status, COUNT(*) as c FROM attendance WHERE student_id=? GROUP BY status",
            (student_id,))
        st = {r["status"]: r["c"] for r in cursor.fetchall()}
        total = sum(st.values())
        if total == 0:
            return {"student_id": student_id, "student_name": s["name"],
                    "total_days": 0, "attendance_rate": 0, "message": "No records"}
        present = st.get("Present", 0)
        return {
            "student_id": student_id, "student_name": s["name"],
            "total_days": total, "present": present,
            "absent": st.get("Absent", 0), "late": st.get("Late", 0),
            "excused": st.get("Excused", 0),
            "attendance_rate": round(present / total * 100, 2),
        }


# ============================================================
# NOTIFICATIONS
# ============================================================
@app.get("/notifications")
def get_notifications(user=Depends(require_user), limit: int = 50):
    accessible = get_accessible_students(user)
    ids = [s["id"] for s in accessible]
    if not ids:
        return {"count": 0, "unread_count": 0, "notifications": []}
    with get_db_connection() as conn:
        cursor = conn.cursor()
        ph = ",".join("?" * len(ids))
        cursor.execute(f"""
            SELECT * FROM notifications WHERE student_id IN ({ph})
            ORDER BY created_at DESC LIMIT ?
        """, ids + [limit])
        notifs = [dict(r) for r in cursor.fetchall()]
        cursor.execute(
            f"SELECT COUNT(*) as c FROM notifications WHERE student_id IN ({ph}) AND is_read=0",
            ids)
        unread = cursor.fetchone()["c"]
        return {"count": len(notifs), "unread_count": unread, "notifications": notifs}


@app.put("/notifications/mark-read/{notif_id}")
def mark_read(notif_id: int, user=Depends(require_user)):
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("UPDATE notifications SET is_read=1 WHERE id=?", (notif_id,))
        conn.commit()
    return {"message": "Read"}


@app.put("/notifications/mark-all-read")
def mark_all_read(user=Depends(require_user)):
    accessible = get_accessible_students(user)
    ids = [s["id"] for s in accessible]
    if not ids:
        return {"message": "No notifications"}
    with get_db_connection() as conn:
        cursor = conn.cursor()
        ph = ",".join("?" * len(ids))
        cursor.execute(f"UPDATE notifications SET is_read=1 WHERE student_id IN ({ph})", ids)
        conn.commit()
    return {"message": "All read"}


# ============================================================
# STATS
# ============================================================
@app.get("/stats/overall")
def overall_stats(user=Depends(require_user)):
    students = get_accessible_students(user)
    if not students:
        return {"message": "No data"}
    avgs = [s["average"] for s in students]
    grades = [s["grade"] for s in students]
    dist = {}
    for g in grades:
        dist[g] = dist.get(g, 0) + 1
    return {
        "total_students": len(students),
        "average_of_averages": round(float(np.mean(avgs)), 2),
        "median_average": round(float(np.median(avgs)), 2),
        "highest_average": round(float(max(avgs)), 2),
        "lowest_average": round(float(min(avgs)), 2),
        "grade_distribution": dist,
        "pass_rate": round(sum(1 for s in students if s["grade"] != "F") / len(students) * 100, 2),
        "distinction_rate": round(
            sum(1 for s in students if s["grade"] in ["A", "A+"]) / len(students) * 100, 2),
        "top_performers": sorted(students, key=lambda x: x["average"], reverse=True)[:5],
    }


@app.get("/analytics/dashboard")
def analytics_dashboard(user=Depends(require_user)):
    students = get_accessible_students(user)
    if not students:
        return {"has_data": False, "total_students": 0, "message": "No data"}
    avgs = [s["average"] for s in students]
    grades = [s["grade"] for s in students]
    dist = {}
    for g in grades:
        dist[g] = dist.get(g, 0) + 1
    ranges = {
        "90-100 (A+)": sum(1 for a in avgs if a >= 90),
        "80-89 (A)": sum(1 for a in avgs if 80 <= a < 90),
        "70-79 (B)": sum(1 for a in avgs if 70 <= a < 80),
        "60-69 (C)": sum(1 for a in avgs if 60 <= a < 70),
        "50-59 (D)": sum(1 for a in avgs if 50 <= a < 60),
        "40-49 (E)": sum(1 for a in avgs if 40 <= a < 50),
        "0-39 (F)": sum(1 for a in avgs if a < 40),
    }
    subj_data = {}
    for s in students:
        for i, subj in enumerate(s["subjects"]):
            subj_data.setdefault(subj, []).append(s["marks"][i] if i < len(s["marks"]) else 0)
    subj_analytics = []
    for subj, marks in subj_data.items():
        arr = np.array(marks)
        subj_analytics.append({
            "subject": subj,
            "average": round(float(np.mean(arr)), 2),
            "highest": float(np.max(arr)),
            "lowest": float(np.min(arr)),
            "pass_rate": round(sum(1 for m in marks if m >= 40) / len(marks) * 100, 2),
            "count": len(marks),
        })
    subj_analytics.sort(key=lambda x: x["average"], reverse=True)
    passed = sum(1 for g in grades if g != "F")
    return {
        "has_data": True,
        "total_students": len(students),
        "overall_average": round(float(np.mean(avgs)), 2),
        "overall_median": round(float(np.median(avgs)), 2),
        "overall_std": round(float(np.std(avgs)), 2),
        "highest_average": round(float(max(avgs)), 2),
        "lowest_average": round(float(min(avgs)), 2),
        "passed": passed, "failed": len(students) - passed,
        "pass_rate": round(passed / len(students) * 100, 2),
        "distinction_rate": round(
            sum(1 for g in grades if g in ["A", "A+"]) / len(students) * 100, 2),
        "grade_distribution": dist, "grade_ranges": ranges,
        "subject_analytics": subj_analytics,
        "top_performers": sorted([
            {"id": s["id"], "name": s["name"], "average": s["average"], "grade": s["grade"]}
            for s in students
        ], key=lambda x: x["average"], reverse=True)[:10],
        "hardest_subject": subj_analytics[-1] if subj_analytics else None,
        "easiest_subject": subj_analytics[0] if subj_analytics else None,
        "timestamp": datetime.now().isoformat(),
    }


@app.get("/trends/{student_id}")
def trends(student_id: int, user=Depends(require_user)):
    if not any(s["id"] == student_id for s in get_accessible_students(user)):
        raise HTTPException(403, "Access denied")
    return analyze_trends(student_id)


# ============================================================
# PDF REPORT CARDS
# ============================================================
@app.get("/students/{student_id}/report-card")
def report_card(student_id: int, user=Depends(require_user)):
    if generate_report_card is None:
        raise HTTPException(500, "PDF not available")
    student = next((s for s in get_accessible_students(user) if s["id"] == student_id), None)
    if not student:
        raise HTTPException(404, "Not found")
    student["recommendations"] = generate_recommendations(student["marks"], student["average"])
    fname = f"report_{student_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    fpath = os.path.join(REPORTS_DIR, fname)
    generate_report_card(student, fpath)
    return FileResponse(fpath, media_type="application/pdf", filename=fname)


@app.get("/students/{student_id}/report-card/{template}")
def report_card_template(
    student_id: int,
    template: str = "modern",
    user=Depends(require_user),
):
    """Generate report card with chosen template."""
    student = next((s for s in get_accessible_students(user) if s["id"] == student_id), None)
    if not student:
        raise HTTPException(404, "Not found")

    student["recommendations"] = generate_recommendations(student["marks"], student["average"])

    fname = f"report_{student_id}_{template}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    fpath = os.path.join(REPORTS_DIR, fname)

    try:
        if template == "modern" and generate_modern_report:
            generate_modern_report(student, fpath)
        elif template == "minimal" and generate_minimal_report:
            generate_minimal_report(student, fpath)
        elif template == "classic":
            if generate_report_card is None:
                raise HTTPException(500, "Classic template unavailable")
            generate_report_card(student, fpath)
        else:
            raise HTTPException(400, f"Unknown template: {template}")

        return FileResponse(fpath, media_type="application/pdf", filename=fname)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"PDF error: {str(e)}")


# ============================================================
# EMAIL QUEUE
# ============================================================
@app.post("/admin/process-email-queue")
def process_queue_manual(admin=Depends(require_admin)):
    process_email_queue()
    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) as c FROM email_queue WHERE sent=0")
        pending = cursor.fetchone()["c"]
    return {"message": "Queue processed", "pending": pending}


# ============================================================
# WEBSOCKET
# ============================================================
@app.websocket("/ws/notifications")
async def websocket_notifications(websocket: WebSocket, token: str = Query("")):
    payload = decode_token(token)
    if not payload:
        await websocket.close(code=4001)
        return

    username = payload.get("sub")
    if not username:
        await websocket.close(code=4001)
        return

    with get_db_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT id FROM users WHERE username=? AND is_active=1", (username,))
        row = cursor.fetchone()
        if not row:
            await websocket.close(code=4001)
            return
        user_id = row["id"]

    await manager.connect(websocket, user_id)

    try:
        await websocket.send_json({
            "type": "connected",
            "message": "Real-time notifications active",
            "user_id": user_id,
        })

        while True:
            try:
                data = await asyncio.wait_for(websocket.receive_text(), timeout=30)
                if data == "ping":
                    await websocket.send_json({"type": "pong"})
            except asyncio.TimeoutError:
                await websocket.send_json({"type": "heartbeat"})
    except WebSocketDisconnect:
        manager.disconnect(websocket, user_id)
        print(f"[WS] User {user_id} disconnected")
    except Exception as e:
        manager.disconnect(websocket, user_id)
        print(f"[WS] Error: {e}")


@app.get("/ws/test")
def ws_test(user=Depends(require_user)):
    """Test endpoint to trigger a live event."""
    log_live_event(user["id"], "test", {"message": "This is a test notification"})

    try:
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        loop.run_until_complete(manager.send_to_user(user["id"], {
            "type": "test",
            "message": "Hello from server!",
            "timestamp": datetime.now().isoformat(),
        }))
        loop.close()
    except Exception as e:
        return {"message": f"Event logged but WS send failed: {e}"}

    return {"message": "Test event sent"}


# ============================================================
# STATIC FILES
# ============================================================
try:
    app.mount("/uploads", StaticFiles(directory=UPLOAD_DIR), name="uploads")
    print(f"[OK] Static files mounted at /uploads")
except Exception as e:
    print(f"WARNING: Could not mount uploads: {e}")


# ============================================================
# BACKGROUND SCHEDULER
# ============================================================
scheduler = BackgroundScheduler()


def auto_backup():
    try:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_filename = f"auto_backup_{timestamp}.zip"
        backup_path = os.path.join(BACKUP_DIR, "auto", backup_filename)

        with zipfile.ZipFile(backup_path, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.write(DB_PATH, "student_marks.db")

        print(f"[AUTO-BACKUP] {backup_filename}")

        auto_dir = os.path.join(BACKUP_DIR, "auto")
        files = sorted(
            [f for f in os.listdir(auto_dir) if f.endswith(".zip")],
            reverse=True,
        )
        for old in files[7:]:
            try:
                os.remove(os.path.join(auto_dir, old))
            except Exception:
                pass
    except Exception as e:
        print(f"[AUTO-BACKUP ERROR] {e}")


scheduler.add_job(auto_backup, IntervalTrigger(hours=AUTO_BACKUP_HOURS), id="auto_backup")
scheduler.add_job(send_weekly_report, IntervalTrigger(days=7), id="weekly_report")


@app.on_event("startup")
def startup_event():
    try:
        scheduler.start()
        print("[OK] Background scheduler started")
    except Exception as e:
        print(f"[WARN] Scheduler start failed: {e}")


@app.on_event("shutdown")
def shutdown_event():
    try:
        scheduler.shutdown(wait=False)
        print("[OK] Scheduler stopped")
    except Exception:
        pass


# ============================================================
# HEALTH
# ============================================================
@app.get("/")
def home():
    return {
        "message": "Student Marks Analyzer API",
        "version": "11.0.0",
        "features": [
            "OTP email verification",
            "username/email login",
            "2FA (TOTP)",
            "parent-child linking",
            "bulk operations",
            "class/cohort analytics",
            "student profiles",
            "Excel/CSV/JSON export",
            "PDF report cards",
            "email notifications",
            "multi-language (EN/HI)",
            "audit logs",
            "role history",
            "backup admin",
            "self-demotion protection",
            "student photo upload",
            "exams & timetable",
            "assignments tracker",
            "fee management",
            "saved filters",
            "database backup & restore",
            "scheduled reports",
            "whatsapp OTP",
            "websocket notifications",
            "customizable PDF templates",
            "attendance heatmap",
        ],
    }

@app.get("/favicon.ico")
def favicon():
    return Response(status_code=204)

@app.get("/health")
def health():
    try:
        with get_db_connection() as conn:
            conn.execute("SELECT 1")
        db = "healthy"
    except Exception:
        db = "unhealthy"
    return {"status": "healthy" if db == "healthy" else "degraded",
            "database": db, "timestamp": datetime.now().isoformat()}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
